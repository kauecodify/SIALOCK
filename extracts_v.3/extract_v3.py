# =====================================================================
#  EXTRACT V.3  ·  Pipeline de Extracao, Enriquecimento, Risco e Treino Neural
#  V.3
# =====================================================================
from __future__ import annotations

import os
import re
import sys
import json
import time
import shutil
import hashlib
import warnings
import threading
import itertools
import argparse
import random
from pathlib import Path
from collections import Counter
from datetime import datetime
from dataclasses import dataclass, field, asdict
from typing import Optional

import numpy as np
import pandas as pd
import joblib
import requests

try:
    import pyautogui
    _HAS_GUI = True
except Exception:
    _HAS_GUI = False

from tenacity import retry, stop_after_attempt, wait_exponential

from sklearn.base import BaseEstimator, ClassifierMixin, clone
from sklearn.compose import ColumnTransformer
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler, OneHotEncoder, Normalizer
from sklearn.impute import SimpleImputer
from sklearn.feature_extraction.text import HashingVectorizer, CountVectorizer
from sklearn.linear_model import SGDClassifier, PassiveAggressiveClassifier
from sklearn.naive_bayes import ComplementNB
from sklearn.svm import LinearSVC
from sklearn.calibration import CalibratedClassifierCV
from sklearn.model_selection import cross_val_score, train_test_split
from sklearn.metrics import (accuracy_score, f1_score, precision_score,
                             recall_score, top_k_accuracy_score, log_loss)
from sklearn.neighbors import LocalOutlierFactor
from sklearn.neural_network import MLPClassifier
from sklearn.decomposition import TruncatedSVD

warnings.filterwarnings("ignore")

# =====================================================================
#  CONSTANTES
# =====================================================================
APP_NOME   = "EXTRACT V.3"
APP_SUB    = "Pipeline de Extracao, Enriquecimento, Risco e Rede Neural"
APP_VERSAO = "3.1"

PASTA_DADOS = "./dados"
PASTA_CACHE = "./cache_externo"
PASTA_SAIDA = "./saida"
PASTA_HIST  = "./historico"
PASTA_PARTS = "./saida/partes"
PASTA_LOGS  = "./logs"
MODELO_PATH = "./modelo_extract_v3.joblib"
HIST_PATH   = os.path.join(PASTA_HIST, "extract_v3_history.json")
PADRAO_PATH = os.path.join(PASTA_HIST, "extract_v3_patterns.json")

CAMPO_ID_DEFAULT       = "CNPJ"
CHUNK_SIZE             = 50_000
SAMPLE_ROWS            = 30_000
HASH_WORD_FEATURES     = 2 ** 18
HASH_CHAR_FEATURES     = 2 ** 17
MAX_CNPJ_EXTERNO       = 500
RANDOM_STATE           = 42
N_SESSOES_PADRAO       = 10
TIMEOUT_API            = 15
TOP_K                  = [1, 3, 5, 10]
N_TOP_RISCO_XLSX       = 5000
LIMIAR_SURGE           = 3.0
LIMIAR_ETA_DIAS        = 7
MAX_CLASSES            = 20000

# --- neural ---
NEURAL_MAX_ROWS        = 8000
NEURAL_HIDDEN          = (256, 128)
NEURAL_EPOCHS          = 15
NEURAL_LR              = 1e-3
NEURAL_BATCH           = 128
NEURAL_SVD_DIM         = 256

PESO_KPIS = {
    "subvalorizacao": 0.20, "disparidade_hs": 0.12,
    "surge_volume": 0.12, "desvio_eta": 0.08,
    "completude": 0.04, "outlier": 0.05,
    "consist_cadastral": 0.13, "idade_empresa": 0.08,
    "score_credito": 0.10, "padrao_char": 0.08,
}

COLUNAS_KPI = [
    "KPI_Subvalorizacao", "KPI_DisparidadeHS", "KPI_SurgeVolume",
    "KPI_DesvioETA", "KPI_Completude", "KPI_Outlier",
    "KPI_Ext_ConsistenciaCadastral", "KPI_Ext_IdadeEmpresa",
    "KPI_Ext_ScoreCredito", "KPI_PadraoChar", "RiskScore",
]

API_BRASILAPI     = "https://brasilapi.com.br/api/cnpj/v1/{cnpj}"
API_MINHA_RECEITA = "https://minhareceita.org/{cnpj}"
API_RECEITAWS     = "https://receitaws.com.br/v1/cnpj/{cnpj}"
API_CNPJWS        = "https://publica.cnpjws.com/cnpj/{cnpj}"

# =====================================================================
#  CONFIG
# =====================================================================
@dataclass
class Config:
    headless:        bool = False
    modo:            str  = "rapido"
    dry_run:         bool = False
    sem_refinamento: bool = False
    sem_xlsx:        bool = False

    pasta_dados:  str = PASTA_DADOS
    pasta_saida:  str = PASTA_SAIDA
    pasta_cache:  str = PASTA_CACHE
    pasta_hist:   str = PASTA_HIST
    pasta_parts:  str = PASTA_PARTS
    pasta_logs:   str = PASTA_LOGS
    modelo_path:  str = MODELO_PATH

    epochs:      int = 0
    sample_rows: int = SAMPLE_ROWS
    max_classes: int = MAX_CLASSES
    max_cnpj:    int = MAX_CNPJ_EXTERNO
    n_top_risco: int = N_TOP_RISCO_XLSX

    log_level: str  = "INFO"
    log_jsonl: bool = True
    no_color:  bool = False

    # --- neural ---
    neural:           bool  = True
    neural_epochs:    int   = NEURAL_EPOCHS
    neural_hidden:    str   = "256,128"
    neural_lr:        float = NEURAL_LR
    neural_max_rows:  int   = NEURAL_MAX_ROWS
    neural_svd:       int   = NEURAL_SVD_DIM
    neural_show_code: bool  = True
    neural_delay:     float = 0.015

    # --- stream enriquecido ---
    stream_enriquecido: bool = False
    stream_budget:      int  = 500
    stream_so_cache:    bool = False

    # --- colunas forçadas ---
    colunas_texto: list = field(default_factory=list)
    colunas_cat:   list = field(default_factory=list)
    colunas_num:   list = field(default_factory=list)

    def aplicar(self):
        global PASTA_DADOS, PASTA_SAIDA, PASTA_CACHE, PASTA_HIST
        global PASTA_PARTS, PASTA_LOGS, MODELO_PATH
        PASTA_DADOS = self.pasta_dados
        PASTA_SAIDA = self.pasta_saida
        PASTA_CACHE = self.pasta_cache
        PASTA_HIST  = self.pasta_hist
        PASTA_PARTS = self.pasta_parts
        PASTA_LOGS  = self.pasta_logs
        MODELO_PATH = self.modelo_path

    def hidden_tuple(self) -> tuple:
        try:
            return tuple(int(x) for x in str(self.neural_hidden).split(",") if x.strip())
        except Exception:
            return NEURAL_HIDDEN

    def to_dict(self) -> dict:
        d = asdict(self)
        for k in ("colunas_texto", "colunas_cat", "colunas_num"):
            d.pop(k, None)
        return d


def carregar_config_arquivo(caminho: str) -> dict:
    if not caminho or not os.path.exists(caminho):
        return {}
    try:
        with open(caminho, "r", encoding="utf-8") as f:
            return json.load(f) or {}
    except Exception as e:
        print(f"[WARN] config {caminho}: {e}")
        return {}


def parse_args(argv=None) -> Config:
    p = argparse.ArgumentParser(
        prog="extract-v3",
        description=f"{APP_NOME} · {APP_SUB} — v{APP_VERSAO}",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Exemplos:\n"
            "  extract-v3\n"
            "  extract-v3 --headless --modo rapido --epochs 3\n"
            "  extract-v3 --headless --modo completo --epochs 5 --neural-epochs 20\n"
            "  extract-v3 --headless --config config.json --no-neural\n"
            "  extract-v3 --headless --stream-enriquecido --stream-budget 500\n"
            "  extract-v3 --headless --dry-run\n"
        ),
    )
    p.add_argument("--headless", action="store_true")
    p.add_argument("--modo", choices=["completo", "rapido"], default="rapido")
    p.add_argument("--epochs", type=int, default=0)
    p.add_argument("--config", type=str, default="")
    p.add_argument("--dados", type=str, default=None)
    p.add_argument("--saida", type=str, default=None)
    p.add_argument("--modelo", type=str, default=None)
    p.add_argument("--sample-rows", type=int, default=None)
    p.add_argument("--max-classes", type=int, default=None)
    p.add_argument("--max-cnpj", type=int, default=None)
    p.add_argument("--top-risco", type=int, default=None)
    p.add_argument("--log-level", choices=["DEBUG", "INFO", "WARN", "ERRO"], default="INFO")
    p.add_argument("--no-jsonl", action="store_true")
    p.add_argument("--no-color", action="store_true")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--sem-refinamento", action="store_true")
    p.add_argument("--sem-xlsx", action="store_true")
    p.add_argument("--texto", nargs="*", default=None)
    p.add_argument("--cat", nargs="*", default=None)
    p.add_argument("--num", nargs="*", default=None)

    # neural
    p.add_argument("--no-neural", action="store_true")
    p.add_argument("--neural-epochs", type=int, default=None)
    p.add_argument("--neural-hidden", type=str, default=None)
    p.add_argument("--neural-lr", type=float, default=None)
    p.add_argument("--neural-max-rows", type=int, default=None)
    p.add_argument("--neural-svd", type=int, default=None)
    p.add_argument("--no-neural-viewer", action="store_true")
    p.add_argument("--neural-delay", type=float, default=None)

    # stream enriquecido
    p.add_argument("--stream-enriquecido", action="store_true",
                   help="enriquece CNPJs durante o partial_fit")
    p.add_argument("--stream-budget", type=int, default=None,
                   help="max chamadas API durante o treino (default 500)")
    p.add_argument("--stream-so-cache", action="store_true",
                   help="com --stream-enriquecido, usa apenas cache (zero API)")

    ns = p.parse_args(argv)
    cfg = Config()
    if ns.config:
        for k, v in carregar_config_arquivo(ns.config).items():
            if hasattr(cfg, k):
                setattr(cfg, k, v)

    cfg.headless        = ns.headless
    cfg.modo            = ns.modo
    cfg.dry_run         = ns.dry_run
    cfg.sem_refinamento = ns.sem_refinamento
    cfg.sem_xlsx        = ns.sem_xlsx
    cfg.log_level       = ns.log_level
    cfg.log_jsonl       = not ns.no_jsonl
    cfg.no_color        = ns.no_color
    cfg.neural          = not ns.no_neural
    cfg.neural_show_code = not ns.no_neural_viewer
    cfg.stream_enriquecido = ns.stream_enriquecido
    cfg.stream_so_cache    = ns.stream_so_cache

    if ns.epochs:               cfg.epochs          = ns.epochs
    if ns.dados:                cfg.pasta_dados     = ns.dados
    if ns.saida:                cfg.pasta_saida     = ns.saida
    if ns.modelo:               cfg.modelo_path     = ns.modelo
    if ns.sample_rows:          cfg.sample_rows     = ns.sample_rows
    if ns.max_classes:          cfg.max_classes     = ns.max_classes
    if ns.max_cnpj:             cfg.max_cnpj        = ns.max_cnpj
    if ns.top_risco:            cfg.n_top_risco     = ns.top_risco
    if ns.texto is not None:    cfg.colunas_texto   = ns.texto
    if ns.cat is not None:      cfg.colunas_cat     = ns.cat
    if ns.num is not None:      cfg.colunas_num     = ns.num
    if ns.neural_epochs:        cfg.neural_epochs   = ns.neural_epochs
    if ns.neural_hidden:        cfg.neural_hidden   = ns.neural_hidden
    if ns.neural_lr:            cfg.neural_lr       = ns.neural_lr
    if ns.neural_max_rows:      cfg.neural_max_rows = ns.neural_max_rows
    if ns.neural_svd is not None: cfg.neural_svd    = ns.neural_svd
    if ns.neural_delay is not None: cfg.neural_delay = ns.neural_delay
    if ns.stream_budget:        cfg.stream_budget   = ns.stream_budget

    cfg.aplicar()
    return cfg


# =====================================================================
#  UI
# =====================================================================
class UI:
    RESET   = "\033[0m"
    BOLD    = "\033[1m"
    DIM     = "\033[2m"

    RED     = "\033[38;5;203m"
    GREEN   = "\033[38;5;114m"
    YELLOW  = "\033[38;5;221m"
    BLUE    = "\033[38;5;75m"
    MAGENTA = "\033[38;5;176m"
    CYAN    = "\033[38;5;80m"
    WHITE   = "\033[38;5;255m"
    GRAY    = "\033[38;5;245m"
    ORANGE  = "\033[38;5;215m"

    _CORES_ATIVAS = True
    _jsonl_path: str | None = None
    _jsonl_lock = threading.Lock()

    @classmethod
    def _c(cls, cor: str) -> str:
        return cor if cls._CORES_ATIVAS else ""

    @classmethod
    def c(cls, txt, *cores):
        return "".join(cls._c(c) for c in cores) + str(txt) + cls._c(cls.RESET)

    @classmethod
    def init(cls):
        if os.environ.get("NO_COLOR") or not sys.stdout.isatty():
            cls._CORES_ATIVAS = False
        if os.name == "nt":
            try: os.system("")
            except Exception: pass

    @classmethod
    def init_jsonl(cls, pasta_logs: str, habilitado: bool = True):
        if not habilitado:
            cls._jsonl_path = None; return
        try:
            os.makedirs(pasta_logs, exist_ok=True)
            cls._jsonl_path = os.path.join(
                pasta_logs, f"extract_v3_{datetime.now():%Y%m%d}.jsonl")
        except Exception:
            cls._jsonl_path = None

    @classmethod
    def jsonl(cls, event: str, level: str = "INFO", **campos):
        if not cls._jsonl_path: return
        rec = {"ts": datetime.now().isoformat(timespec="milliseconds"),
               "app": APP_NOME, "ver": APP_VERSAO,
               "level": level, "event": event}
        for k, v in campos.items():
            try:
                if isinstance(v, np.integer):     v = int(v)
                elif isinstance(v, np.floating):  v = float(v)
                elif isinstance(v, np.ndarray):   v = v.tolist()
                elif isinstance(v, pd.Timestamp): v = v.isoformat()
                rec[k] = v
            except Exception:
                rec[k] = str(v)
        with cls._jsonl_lock:
            try:
                with open(cls._jsonl_path, "a", encoding="utf-8") as f:
                    f.write(json.dumps(rec, ensure_ascii=False, default=str) + "\n")
            except Exception:
                pass

    @classmethod
    def banner(cls, titulo: str, sub: str = "", largura: int = 78):
        topo = "╔" + "═" * (largura - 2) + "╗"
        base = "╚" + "═" * (largura - 2) + "╝"
        print(cls.c(topo, cls.CYAN, cls.BOLD))
        print(cls.c("║", cls.CYAN) + cls.c(f"  {titulo}".ljust(largura - 2) + "║", cls.WHITE, cls.BOLD))
        if sub:
            print(cls.c("║", cls.CYAN) + cls.c(f"  {sub}".ljust(largura - 2) + "║", cls.GRAY))
        print(cls.c(base, cls.CYAN, cls.BOLD))

    @classmethod
    def secao(cls, num: str, titulo: str, largura: int = 78):
        print()
        print(cls.c(f"┌─ [{num}] ", cls.MAGENTA, cls.BOLD)
              + cls.c(titulo, cls.WHITE, cls.BOLD)
              + cls.c(" " + "─" * max(0, largura - len(num) - len(titulo) - 8), cls.MAGENTA))

    @classmethod
    def fim_secao(cls):
        print(cls.c("└" + "─" * 76, cls.MAGENTA))

    @classmethod
    def log(cls, nivel: str, msg: str):
        ts = datetime.now().strftime("%H:%M:%S")
        mapa = {"INFO": (cls.BLUE, "•"), "OK": (cls.GREEN, "✔"),
                "WARN": (cls.YELLOW, "▲"), "ERRO": (cls.RED, "✖"),
                "DBG": (cls.GRAY, "…"), "ETAPA": (cls.MAGENTA, "▸")}
        cor, ico = mapa.get(nivel, (cls.WHITE, " "))
        print(f"{cls.c(ts, cls.GRAY)}  {cls.c(ico, cor, cls.BOLD)} "
              f"{cls.c(msg, cls.WHITE if nivel in ('OK','INFO') else cor)}")
        cls.jsonl("log", level=nivel, msg=msg)

    @classmethod
    def info(cls, m):  cls.log("INFO", m)
    @classmethod
    def ok(cls, m):    cls.log("OK", m)
    @classmethod
    def warn(cls, m):  cls.log("WARN", m)
    @classmethod
    def erro(cls, m):  cls.log("ERRO", m)
    @classmethod
    def etapa(cls, m): cls.log("ETAPA", m)

    @classmethod
    def progresso(cls, atual, total, prefixo="", largura=32, extra=""):
        if total <= 0: return
        pct = min(1.0, atual / total)
        cheios = int(pct * largura)
        barra = cls.c("█" * cheios, cls.GREEN) + cls.c("░" * (largura - cheios), cls.GRAY)
        print(f"\r  {cls.c(prefixo, cls.CYAN)} [{barra}] "
              f"{cls.c(f'{pct*100:5.1f}%', cls.WHITE, cls.BOLD)} "
              f"{cls.c(extra, cls.GRAY)}", end="", flush=True)

    @classmethod
    def fim_progresso(cls, msg: str = ""):
        print()
        if msg: cls.ok(msg)

    class Spinner:
        def __init__(self, msg: str):
            self.msg = msg
            self.stop = threading.Event()
            self.t = threading.Thread(target=self._run, daemon=True)
        def _run(self):
            for ch in itertools.cycle("⠋⠙⠹⠸⠼⠴⠦⠧⠇⠏"):
                if self.stop.is_set(): break
                print(f"\r  {UI.c(ch, UI.CYAN, UI.BOLD)} {UI.c(self.msg, UI.WHITE)}   ",
                      end="", flush=True)
                time.sleep(0.08)
        def __enter__(self): self.t.start(); return self
        def __exit__(self, *a):
            self.stop.set(); self.t.join(timeout=0.3)
            print("\r" + " " * (len(self.msg) + 10) + "\r", end="")

    @classmethod
    def tabela(cls, cabecalhos, linhas, larguras=None):
        if not larguras:
            larguras = [max([len(str(c))] + [len(str(r[i])) for r in linhas])
                        if linhas else len(str(c))
                        for i, c in enumerate(cabecalhos)]
        sep = cls.c("  " + "─" * (sum(larguras) + 3 * len(larguras)), cls.GRAY)
        print("  " + "   ".join(cls.c(str(h).ljust(larguras[i]), cls.CYAN, cls.BOLD)
                                for i, h in enumerate(cabecalhos)))
        print(sep)
        for r in linhas:
            print("  " + "   ".join(cls.c(str(v).ljust(larguras[i]), cls.WHITE)
                                    for i, v in enumerate(r)))
        print(sep)

    @classmethod
    def caixa(cls, titulo: str, conteudo: str, cor=None, largura: int = 78):
        cor = cor or cls.CYAN
        print(cls.c("┌" + "─" * (largura - 2) + "┐", cor))
        if titulo:
            print(cls.c("│ ", cor) + cls.c(titulo.ljust(largura - 4), cls.WHITE, cls.BOLD)
                  + cls.c(" │", cor))
            print(cls.c("├" + "─" * (largura - 2) + "┤", cor))
        for l in (conteudo.splitlines() or [""]):
            print(cls.c("│ ", cor) + cls.c(l.ljust(largura - 4), cls.WHITE) + cls.c(" │", cor))
        print(cls.c("└" + "─" * (largura - 2) + "┘", cor))

    @classmethod
    def prompt(cls, msg: str, default: str = "") -> str:
        dica = f" {cls.c('['+default+']', cls.GRAY)}" if default else ""
        try:
            return input(f"  {cls.c('›', cls.GREEN, cls.BOLD)} {cls.c(msg, cls.WHITE)}{dica}: ").strip() or default
        except EOFError:
            return default

    @classmethod
    def menu(cls, titulo: str, opcoes: list) -> str:
        print()
        print(cls.c(f"  {titulo}", cls.WHITE, cls.BOLD))
        for k, d in opcoes:
            print(f"    {cls.c('['+k+']', cls.CYAN, cls.BOLD)} {cls.c(d, cls.WHITE)}")
        try:
            return input(f"  {cls.c('›', cls.GREEN, cls.BOLD)} ").strip().lower()
        except EOFError:
            return opcoes[-1][0]


def log_jsonl(event: str, level: str = "INFO", **campos):
    UI.jsonl(event, level=level, **campos)


# =====================================================================
#  CODEVIEWER
# =====================================================================
class CodeViewer:
    def __init__(self, titulo: str = "neural_train.py", max_vis: int = 16,
                 delay: float = 0.02, enabled: bool = True):
        self.titulo = titulo
        self.max_vis = max_vis
        self.delay = delay
        self.enabled = enabled and UI._CORES_ATIVAS
        self.buffer: list[str] = []
        self._printed = 0
        self._lock = threading.Lock()
        self._aberto = False

    def _render(self) -> list[str]:
        linhas = []
        L = 76
        topo = "┌─ " + self.titulo + " " + "─" * max(0, L - len(self.titulo) - 4)
        linhas.append(UI.c(topo, UI.GRAY))
        start = max(0, len(self.buffer) - self.max_vis)
        for i, l in enumerate(self.buffer[start:], start=start):
            num = f"{i+1:>3} "
            cur = (i == len(self.buffer) - 1)
            if cur:
                linhas.append(UI.c(num, UI.YELLOW, UI.BOLD)
                              + UI.c("▶ ", UI.YELLOW)
                              + UI.c(l, UI.WHITE, UI.BOLD))
            else:
                linhas.append(UI.c(num, UI.GRAY)
                              + UI.c("  ", UI.GRAY)
                              + UI.c(l, UI.GRAY))
        linhas.append(UI.c("└" + "─" * L, UI.GRAY))
        return linhas

    def _redraw(self):
        if not self.enabled: return
        if self._printed:
            sys.stdout.write(f"\033[{self._printed}F")
        linhas = self._render()
        for l in linhas:
            sys.stdout.write("\033[K" + l + "\n")
        self._printed = len(linhas)
        sys.stdout.flush()

    def abrir(self):
        if not self.enabled: return
        self._aberto = True
        self._printed = 0
        self._redraw()

    def escrever(self, linha: str, delay: Optional[float] = None):
        with self._lock:
            if not self.enabled:
                print(f"  {linha}")
                return
            if not self._aberto: self.abrir()
            if delay is None: delay = self.delay
            if delay > 0 and len(linha) > 0:
                parcial = ""
                for ch in linha:
                    parcial += ch
                    if len(self.buffer) == 0 or self.buffer[-1] != parcial:
                        if self.buffer and self.buffer[-1].startswith("▶ "):
                            self.buffer.pop()
                        self.buffer.append(parcial)
                    self._redraw()
                    time.sleep(delay / max(1, len(linha)) * 4)
            else:
                self.buffer.append(linha)
                self._redraw()
            UI.jsonl("code_line", line=linha)

    def bloco(self, linhas: list[str], delay: Optional[float] = None):
        for l in linhas:
            self.escrever(l, delay=delay)

    def pausa(self, s: float = 0.15):
        if self.enabled: time.sleep(s)

    def fim(self, msg: str = ""):
        if not self.enabled: return
        if msg: self.escrever(f"# ✔ {msg}", delay=0)
        self._redraw()
        print()
        self._aberto = False


# =====================================================================
#  LIVETABLE
# =====================================================================
class LiveTable:
    def __init__(self, colunas: list[str], titulo: str = "",
                 larguras: Optional[list[int]] = None, enabled: bool = True):
        self.colunas = colunas
        self.titulo = titulo
        self.enabled = enabled and UI._CORES_ATIVAS
        self.larguras = larguras or self._auto_larguras(colunas)
        self.rows: list[list] = []
        self._printed = 0
        self._lock = threading.Lock()
        self._aberto = False

    @staticmethod
    def _auto_larguras(cols) -> list[int]:
        return [max(9, len(str(c)) + 2) for c in cols]

    def _formatar(self, v) -> str:
        if isinstance(v, float):
            if abs(v) < 1e-3 and v != 0: return f"{v:.6f}"
            return f"{v:.4f}"
        if isinstance(v, int):
            return f"{v:,}".replace(",", ".")
        return str(v)

    def _render(self) -> list[str]:
        L = sum(self.larguras) + 3 * (len(self.larguras) - 1) + 4
        linhas = []
        if self.titulo:
            t = "┌─ " + self.titulo + " " + "─" * max(0, L - len(self.titulo) - 4)
            linhas.append(UI.c(t, UI.CYAN, UI.BOLD))
        else:
            linhas.append(UI.c("┌" + "─" * L, UI.CYAN))
        head = "│ " + " │ ".join(
            UI.c(str(c).ljust(self.larguras[i]), UI.CYAN, UI.BOLD)
            for i, c in enumerate(self.colunas)) + " │"
        linhas.append(head)
        sep = "├─" + "─┼─".join("─" * w for w in self.larguras) + "─┤"
        linhas.append(UI.c(sep, UI.GRAY))
        for r in self.rows:
            linha = "│ " + " │ ".join(
                UI.c(self._formatar(v).ljust(self.larguras[i]), UI.WHITE)
                for i, v in enumerate(r)) + " │"
            linhas.append(linha)
        linhas.append(UI.c("└─" + "─┴─".join("─" * w for w in self.larguras) + "─┘",
                           UI.CYAN))
        return linhas

    def _redraw(self):
        if not self.enabled: return
        if self._printed:
            sys.stdout.write(f"\033[{self._printed}F")
        linhas = self._render()
        for l in linhas:
            sys.stdout.write("\033[K" + l + "\n")
        self._printed = len(linhas)
        sys.stdout.flush()

    def abrir(self):
        if not self.enabled: return
        self._aberto = True
        self._printed = 0
        self._redraw()

    def add(self, row: list):
        with self._lock:
            if not self.enabled:
                print("  " + " | ".join(str(x) for x in row))
                UI.jsonl("table_row", row=row)
                return
            if not self._aberto: self.abrir()
            self.rows.append(row)
            self._redraw()
            UI.jsonl("table_row", row=[self._formatar(v) for v in row])

    def atualizar_ultima(self, row: list):
        with self._lock:
            if not self.enabled: return
            if self.rows:
                self.rows[-1] = row
                self._redraw()

    def fim(self, msg: str = ""):
        if not self.enabled: return
        self._redraw()
        if msg:
            print("  " + UI.c(msg, UI.GREEN, UI.BOLD))
        print()
        self._aberto = False


# =====================================================================
#  HELPERS DE UI (GUI fallback)
# =====================================================================
def alertar(msg: str, titulo: str = "Extract V.3"):
    if _HAS_GUI:
        try:
            pyautogui.alert(msg, titulo, "OK"); return
        except Exception:
            pass
    UI.caixa(titulo, msg, cor=UI.GREEN if "erro" not in titulo.lower() else UI.RED)


def menu_inicial() -> str:
    UI.banner(f"{APP_NOME}  ·  v{APP_VERSAO}", APP_SUB)
    print()
    op = UI.menu("Como deseja executar?", [
        ("1", "Completo  —  enriquece APIs + treino continuo + refinamento + rede neural"),
        ("2", "Rapido    —  apenas dados locais (sem APIs externas)"),
        ("3", "Config    —  ajustar parametros"),
        ("0", "Cancelar"),
    ])
    return {"1": "Completo", "2": "Rapido", "3": "Config"}.get(op, "Cancelar")


def captura_clipboard(titulo="Captura de PDF") -> str | None:
    if _HAS_GUI:
        try:
            resp = pyautogui.confirm(
                "Copie o conteudo do PDF (Ctrl+C) e clique em 'Ja copiei'.",
                titulo, ["Ja copiei", "Cancelar"])
            if resp != "Ja copiei": return None
            import pyperclip
            return pyperclip.paste()
        except Exception:
            pass
    UI.info("Cole o texto e pressione Enter duas vezes:")
    linhas = []
    try:
        while True:
            l = input()
            if not l: break
            linhas.append(l)
    except EOFError:
        pass
    return "\n".join(linhas) or None


def coletar_valores_alvo(colunas) -> dict | None:
    valores = {}
    for col in colunas:
        try:
            v = UI.prompt(f"Valor para '{col}'")
        except Exception:
            return None
        if v is None: return None
        valores[col] = v
    return valores


# =====================================================================
#  HELPERS DE DADOS
# =====================================================================
def listar_arquivos(pasta: str):
    if not os.path.isdir(pasta): return []
    exts = {".xlsx", ".parquet", ".csv", ".pdf", ".txt"}
    return [p for p in sorted(Path(pasta).glob("*")) if p.suffix.lower() in exts]


def hash_cabecalho(df: pd.DataFrame) -> str:
    return hashlib.md5(",".join(map(str, df.columns)).encode()).hexdigest()[:10]


def normalizar_id(serie: pd.Series) -> pd.Series:
    return serie.astype(str).str.strip().str.replace(r"\s+", "", regex=True)


def normalizar_id_cnpj(serie: pd.Series) -> pd.Series:
    return serie.astype(str).str.replace(r"\D", "", regex=True).str.zfill(14)


def extrair_texto_pdf(caminho: str) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(caminho) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception:
        pass
    try:
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(caminho).pages)
    except Exception:
        pass
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(caminho), lang="por")
    except Exception:
        return ""


def iterar_chunks(pasta: str, chunk_size: int = CHUNK_SIZE):
    for arq in listar_arquivos(pasta):
        suf = arq.suffix.lower()
        try:
            if suf == ".csv":
                for chunk in pd.read_csv(arq, chunksize=chunk_size, dtype=str,
                                         sep=None, engine="python"):
                    chunk.columns = [str(c).strip() for c in chunk.columns]
                    yield arq.name, chunk
            elif suf == ".xlsx":
                try:
                    df = pd.read_excel(arq, dtype=str)
                except Exception:
                    continue
                df.columns = [str(c).strip() for c in df.columns]
                for i in range(0, len(df), chunk_size):
                    yield arq.name, df.iloc[i:i + chunk_size].copy()
            elif suf == ".parquet":
                try:
                    import pyarrow.parquet as pq
                    pf = pq.ParquetFile(arq)
                    for batch in pf.iter_batches(batch_size=chunk_size):
                        df = batch.to_pandas()
                        df.columns = [str(c).strip() for c in df.columns]
                        yield arq.name, df
                except Exception:
                    df = pd.read_parquet(arq, engine="pyarrow")
                    df.columns = [str(c).strip() for c in df.columns]
                    for i in range(0, len(df), chunk_size):
                        yield arq.name, df.iloc[i:i + chunk_size].copy()
            elif suf in (".pdf", ".txt"):
                txt = extrair_texto_pdf(str(arq)) if suf == ".pdf" \
                    else arq.read_text(encoding="utf-8", errors="ignore")
                if not txt.strip(): continue
                fmt = re.findall(r"\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}", txt)
                raw = re.findall(r"\b\d{14}\b", txt)
                todos = list(dict.fromkeys(
                    [re.sub(r"\D", "", c).zfill(14) for c in fmt] +
                    [c.zfill(14) for c in raw]))
                if not todos: continue
                for i in range(0, len(todos), chunk_size):
                    sub = todos[i:i + chunk_size]
                    yield arq.name, pd.DataFrame({
                        CAMPO_ID_DEFAULT: sub,
                        "TextoPDF": [txt[:5000]] * len(sub)})
        except Exception as e:
            UI.warn(f"{arq.name}: {e}")


def amostrar_dados(pasta: str, n: int = SAMPLE_ROWS) -> pd.DataFrame:
    partes, total = [], 0
    for _, chunk in iterar_chunks(pasta, chunk_size=min(n, CHUNK_SIZE)):
        if total + len(chunk) > n:
            chunk = chunk.head(n - total)
        partes.append(chunk)
        total += len(chunk)
        if total >= n: break
    if not partes: return pd.DataFrame()
    return pd.concat(partes, ignore_index=True)


class Amostrador:
    def __init__(self, max_rows: int, seed: int = 42):
        self.max_rows = max_rows
        self.chunks = []
        self.n_rows = 0
        self.rng = np.random.RandomState(seed)
    def add(self, df: pd.DataFrame):
        if self.n_rows >= self.max_rows or len(df) == 0: return
        if len(df) > self.max_rows - self.n_rows:
            df = df.sample(n=self.max_rows - self.n_rows,
                           random_state=int(self.rng.randint(0, 2**31 - 1)))
        self.chunks.append(df.copy()); self.n_rows += len(df)
    def get(self) -> pd.DataFrame:
        if not self.chunks: return pd.DataFrame()
        return pd.concat(self.chunks, ignore_index=True)


def detectar_coluna_id(df: pd.DataFrame) -> str | None:
    if CAMPO_ID_DEFAULT in df.columns: return CAMPO_ID_DEFAULT
    candidatos = ["cnpj", "cnpj_cpf", "nr_cnpj", "num_cnpj", "numero_ordem",
                  "num_ordem", "nro_ordem", "ordem", "id", "codigo", "cod",
                  "numero", "nro", "num", "processo", "chave"]
    for c in df.columns:
        if str(c).lower().strip() in candidatos: return c
    for c in df.columns:
        s = df[c].dropna().astype(str).head(100)
        if len(s) == 0: continue
        d = s.str.replace(r"\D", "", regex=True)
        if d.str.len().between(6, 20).mean() > 0.8: return c
    return None


def classificar_colunas(df: pd.DataFrame, coluna_id: str) -> dict:
    cols = {"texto": [], "categoricas": [], "numericas": [],
            "temporais": [], "ignorar": []}
    n = len(df)
    for col in df.columns:
        if col == coluna_id:
            cols["ignorar"].append(col); continue
        s = df[col]
        n_u = s.nunique(dropna=True)
        nulos = float(s.isna().mean())
        if nulos > 0.95 or n_u <= 1:
            cols["ignorar"].append(col); continue
        if pd.api.types.is_datetime64_any_dtype(s):
            cols["temporais"].append(col); continue
        amostra = s.dropna().astype(str).head(200)
        d = amostra.str.replace(r"\D", "", regex=True)
        if len(amostra) and d.str.len().between(8, 20).mean() > 0.7 \
                and n_u > 0.5 * n:
            cols["ignorar"].append(col); continue
        num = pd.to_numeric(s, errors="coerce")
        if num.notna().mean() > 0.7:
            (cols["categoricas"] if n_u < 20 else cols["numericas"]).append(col)
            continue
        (cols["categoricas"] if n_u < 50 else cols["texto"]).append(col)
    return cols


def selecionar_colunas_terminal(classificacao: dict, df: pd.DataFrame) -> dict:
    UI.secao("F0", "Selecao de Colunas")
    for tipo, cols in classificacao.items():
        if not cols: continue
        cor = {"texto": UI.CYAN, "categoricas": UI.MAGENTA,
               "numericas": UI.GREEN, "temporais": UI.YELLOW,
               "ignorar": UI.GRAY}.get(tipo, UI.WHITE)
        print()
        print(f"  {UI.c(f'[{tipo.upper()}]', cor, UI.BOLD)} "
              f"{UI.c(f'({len(cols)})', UI.GRAY)}")
        for c in cols[:50]:
            try:
                u = df[c].nunique(dropna=True)
                print(f"    {UI.c('·', cor)} {UI.c(c, UI.WHITE)} "
                      f"{UI.c(f'(unicos: {u})', UI.GRAY)}")
            except Exception:
                print(f"    {UI.c('·', cor)} {UI.c(c, UI.WHITE)}")
        if len(cols) > 50:
            print(f"    {UI.c(f'... +{len(cols)-50} ocultas', UI.GRAY)}")

    sel = {}
    for tipo in ["texto", "categoricas", "numericas", "temporais"]:
        lista = classificacao.get(tipo, [])
        if not lista:
            sel[tipo] = []; continue
        print()
        print(f"  {UI.c(tipo.upper(), UI.BOLD, UI.WHITE)}:")
        for i, c in enumerate(lista, 1):
            print(f"    {UI.c(f'[{i:>2}]', UI.CYAN)} {UI.c(c, UI.WHITE)}")
        print(f"    {UI.c('[0]', UI.GRAY)} nenhuma   "
              f"{UI.c('[all]', UI.GREEN)} todas   "
              f"{UI.c('[enter]', UI.GREEN)} todas")
        e = UI.prompt("Escolha (ex: 1,3,5)").lower()
        if e in ("all", ""):
            sel[tipo] = list(lista)
        elif e == "0":
            sel[tipo] = []
        else:
            try:
                idx = [int(x.strip()) - 1 for x in e.split(",") if x.strip()]
                sel[tipo] = [lista[i] for i in idx if 0 <= i < len(lista)]
            except Exception:
                sel[tipo] = list(lista)
    UI.fim_secao()
    return sel


# =====================================================================
#  APIs DE CNPJ
# =====================================================================
@retry(stop=stop_after_attempt(2), wait=wait_exponential(min=1, max=4))
def _get_json(url: str, timeout: int = TIMEOUT_API):
    r = requests.get(url, timeout=timeout, headers={"User-Agent": "ExtractV3/1.0"})
    if r.status_code == 200: return r.json()
    if r.status_code in (404, 400, 429): return None
    r.raise_for_status()


def consultar_brasilapi(cnpj: str):
    try:
        d = _get_json(API_BRASILAPI.format(cnpj=cnpj))
        if not d: return None
        return {
            "EXT_RazaoSocial": d.get("razao_social"),
            "EXT_NomeFantasia": d.get("nome_fantasia"),
            "EXT_CNAE_Descricao": d.get("cnae_fiscal_descricao"),
            "EXT_NaturezaJuridica": d.get("natureza_juridica"),
            "EXT_UF": d.get("uf"), "EXT_Municipio": d.get("municipio"),
            "EXT_Situacao": d.get("descricao_situacao_cadastral"),
            "EXT_Porte": d.get("porte"),
            "EXT_CapitalSocial": d.get("capital_social"),
            "EXT_DataAbertura": d.get("data_inicio_atividade"),
            "EXT_QtdSocios": len(d.get("qsa") or []),
            "EXT_QtdCNAESecundarios": len(d.get("cnaes_secundarios") or []),
            "EXT_Email": d.get("email"),
            "EXT_Telefone": d.get("ddd_telefone_1"),
        }
    except Exception:
        return None


def consultar_minha_receita(cnpj: str):
    try:
        d = _get_json(API_MINHA_RECEITA.format(cnpj=cnpj))
        if not d: return None
        return {
            "EXT_Simples": "Sim" if d.get("opcao_pelo_simples") else "Nao",
            "EXT_MEI": "Sim" if d.get("opcao_pelo_mei") else "Nao",
            "EXT_MR_Email": d.get("email"),
            "EXT_MR_Telefone": d.get("ddd_telefone_1"),
        }
    except Exception:
        return None


def consultar_receitaws(cnpj: str):
    try:
        d = _get_json(API_RECEITAWS.format(cnpj=cnpj))
        if not d or d.get("status") == "ERROR": return None
        return {
            "EXT_RW_Situacao": d.get("situacao"),
            "EXT_RW_Atividade": d.get("atividade_principal", [{}])[0].get("text"),
        }
    except Exception:
        return None


def consultar_cnpjws(cnpj: str):
    try:
        d = _get_json(API_CNPJWS.format(cnpj=cnpj))
        if not d or "estabelecimento" not in d: return None
        est = d["estabelecimento"]
        return {
            "EXT_WS_Bairro": est.get("bairro"),
            "EXT_WS_CEP": est.get("cep"),
            "EXT_WS_Cidade": (est.get("cidade") or {}).get("nome"),
        }
    except Exception:
        return None


def uf_para_regiao(uf) -> str:
    if not uf: return "Desconhecida"
    uf = str(uf).upper().strip()
    mapa = {
        "AM": "Norte", "RR": "Norte", "AP": "Norte", "PA": "Norte",
        "TO": "Norte", "RO": "Norte", "AC": "Norte",
        "MA": "Nordeste", "PI": "Nordeste", "CE": "Nordeste", "RN": "Nordeste",
        "PB": "Nordeste", "PE": "Nordeste", "AL": "Nordeste", "SE": "Nordeste",
        "BA": "Nordeste", "MT": "Centro-Oeste", "MS": "Centro-Oeste",
        "GO": "Centro-Oeste", "DF": "Centro-Oeste",
        "SP": "Sudeste", "RJ": "Sudeste", "MG": "Sudeste", "ES": "Sudeste",
        "PR": "Sul", "SC": "Sul", "RS": "Sul",
    }
    return mapa.get(uf, "Desconhecida")


def _score_credito_heuristica(reg: dict) -> float:
    score = 0.5
    sit = str(reg.get("EXT_Situacao") or "").lower()
    if "ativa" in sit: score += 0.2
    elif "baixada" in sit or "suspensa" in sit: score -= 0.3
    idade = reg.get("EXT_IdadeEmpresa") or 0
    if idade >= 10: score += 0.15
    elif idade < 2: score -= 0.15
    try:
        cap = float(reg.get("EXT_CapitalSocial") or 0)
        if cap > 1_000_000: score += 0.1
        elif cap < 10_000: score -= 0.05
    except Exception:
        pass
    if (reg.get("EXT_QtdSocios") or 0) > 2: score += 0.05
    return float(np.clip(score, 0, 1))


def enriquecer_cnpjs(df: pd.DataFrame, max_cnpjs: int = MAX_CNPJ_EXTERNO,
                     so_cache: bool = False) -> pd.DataFrame:
    if df.empty: return df
    os.makedirs(PASTA_CACHE, exist_ok=True)
    cache_file = os.path.join(PASTA_CACHE, "cnpj_cache.parquet")
    cache_df = pd.DataFrame()
    if os.path.exists(cache_file):
        try: cache_df = pd.read_parquet(cache_file, engine="pyarrow")
        except Exception: cache_df = pd.DataFrame()
    col_id = CAMPO_ID_DEFAULT if CAMPO_ID_DEFAULT in df.columns else None
    if col_id is None: return df

    df[col_id] = normalizar_id_cnpj(df[col_id])

    if so_cache:
        if not cache_df.empty and col_id in cache_df.columns:
            return df.merge(cache_df, on=col_id, how="left", suffixes=("", "_ext"))
        return df

    cnpjs = df[col_id].dropna().unique().tolist()
    cacheados = set(cache_df[col_id]) if not cache_df.empty and col_id in cache_df.columns else set()
    faltantes = [c for c in cnpjs if c not in cacheados]
    if len(faltantes) > max_cnpjs: faltantes = faltantes[:max_cnpjs]
    if not faltantes:
        if not cache_df.empty and col_id in cache_df.columns:
            return df.merge(cache_df, on=col_id, how="left", suffixes=("", "_ext"))
        return df

    novos = []; total = len(faltantes)
    UI.info(f"Consultando {total} CNPJs em 4 APIs...")
    for i, c in enumerate(faltantes, 1):
        reg = {col_id: c}
        for fn in (consultar_brasilapi, consultar_minha_receita,
                   consultar_receitaws, consultar_cnpjws):
            r = fn(c)
            if r: reg.update(r)
        reg["EXT_Regiao"] = uf_para_regiao(reg.get("EXT_UF"))
        dt = reg.get("EXT_DataAbertura")
        if dt:
            try:
                d = pd.to_datetime(dt, errors="coerce")
                if pd.notna(d):
                    reg["EXT_IdadeEmpresa"] = round(
                        (pd.Timestamp.now() - d).days / 365.25, 2)
            except Exception:
                pass
        reg["EXT_ScoreCredito"] = _score_credito_heuristica(reg)
        novos.append(reg)
        if i % 5 == 0 or i == total:
            UI.progresso(i, total, "APIs", extra=f"{i}/{total}")
        time.sleep(0.1)
    UI.fim_progresso(f"{len(novos)} CNPJs enriquecidos")

    if novos:
        cache_df = pd.concat([cache_df, pd.DataFrame(novos)], ignore_index=True) \
            if not cache_df.empty else pd.DataFrame(novos)
        cache_df = cache_df.drop_duplicates(subset=[col_id], keep="last")
        try: cache_df.to_parquet(cache_file, engine="pyarrow", index=False)
        except Exception: pass

    if not cache_df.empty and col_id in cache_df.columns:
        df = df.merge(cache_df, on=col_id, how="left", suffixes=("", "_ext"))
    return df


# =====================================================================
#  FEATURES / KPIs / RISCO
# =====================================================================
def adicionar_features_avancadas(df: pd.DataFrame, col_txt, col_cat) -> pd.DataFrame:
    df = df.copy()
    for c in col_txt:
        if c in df.columns:
            s = df[c].astype(str)
            df[f"FE_Len_{c}"] = s.str.len().values
            df[f"FE_Words_{c}"] = s.str.split().str.len().fillna(0).values
    for c in col_cat:
        if c in df.columns:
            try:
                freq = df[c].value_counts(normalize=True)
                df[f"FE_Freq_{c}"] = df[c].map(freq).fillna(0).values
            except Exception:
                pass
    return df


def calcular_kpis_negocio(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    val = pd.to_numeric(df["ValorImportado"], errors="coerce") \
        if "ValorImportado" in df.columns else None
    peso = pd.to_numeric(df["PesoKg"], errors="coerce") \
        if "PesoKg" in df.columns else None

    if val is not None and peso is not None:
        vk = val / peso.replace(0, np.nan)
        med = vk.median()
        df["KPI_Subvalorizacao"] = np.clip((med - vk) / (med if med else 1), 0, 1).fillna(0).values
        df["FE_ValorPorKg"] = vk.fillna(0).values
    else:
        df["KPI_Subvalorizacao"] = 0.0; df["FE_ValorPorKg"] = 0.0

    if "HS_Code" in df.columns:
        hs = df["HS_Code"].astype(str).str.replace(r"\D", "", regex=True)
        z = hs.str.count(r"0+$") / hs.str.len().replace(0, 1)
        df["KPI_DisparidadeHS"] = np.clip(z, 0, 1).fillna(0).values
    else:
        df["KPI_DisparidadeHS"] = 0.0

    if "Importador" in df.columns and val is not None:
        df["_v"] = val
        stats = df.groupby("Importador")["_v"].agg(["mean", "std"]).reset_index()
        stats.columns = ["Importador", "media", "std"]
        df = df.merge(stats, on="Importador", how="left")
        df["KPI_SurgeVolume"] = np.clip(
            (val - df["media"]) / (df["std"].replace(0, np.nan) * LIMIAR_SURGE),
            0, 1).fillna(0).values
        df = df.drop(columns=["_v", "media", "std"])
    else:
        df["KPI_SurgeVolume"] = 0.0

    if "ETA" in df.columns and "ETAReal" in df.columns:
        e = pd.to_datetime(df["ETA"], errors="coerce")
        r = pd.to_datetime(df["ETAReal"], errors="coerce")
        d = (r - e).dt.days.abs()
        df["KPI_DesvioETA"] = np.clip(d / LIMIAR_ETA_DIAS, 0, 1).fillna(0).values
    else:
        df["KPI_DesvioETA"] = 0.0

    df["KPI_Completude"] = 1 - df.notna().mean(axis=1)

    if "EXT_Situacao" in df.columns:
        at = df["EXT_Situacao"].astype(str).str.lower().str.contains("ativa")
        df["KPI_Ext_ConsistenciaCadastral"] = (~at).astype(float).values
    else:
        df["KPI_Ext_ConsistenciaCadastral"] = 0.0

    if "EXT_IdadeEmpresa" in df.columns:
        idade = pd.to_numeric(df["EXT_IdadeEmpresa"], errors="coerce")
        df["KPI_Ext_IdadeEmpresa"] = np.clip(1 - idade / 10, 0, 1).fillna(0.5).values
    else:
        df["KPI_Ext_IdadeEmpresa"] = 0.0

    if "EXT_ScoreCredito" in df.columns:
        sc = pd.to_numeric(df["EXT_ScoreCredito"], errors="coerce")
        df["KPI_Ext_ScoreCredito"] = (1 - sc).fillna(0.5).values
    else:
        df["KPI_Ext_ScoreCredito"] = 0.0

    return df


def fit_char_ngram(df: pd.DataFrame, col_txt, top_n: int = 50,
                   ngram_range=(2, 5)) -> tuple:
    if not col_txt: return None, {}
    cols = [c for c in col_txt if c in df.columns]
    if not cols: return None, {}
    blob = df[cols].fillna("").astype(str).agg(" ".join, axis=1)
    blob = blob.str.lower().str.replace(r"[^a-z0-9]", "", regex=True)
    try:
        vec = CountVectorizer(analyzer="char_wb", ngram_range=ngram_range,
                              min_df=2, max_features=300)
        M = vec.fit_transform(blob)
    except ValueError:
        return None, {}
    vocab = vec.get_feature_names_out()
    freq = np.asarray(M.sum(axis=0)).ravel()
    idx = np.argsort(freq)[::-1][:top_n]
    top = {vocab[i]: int(freq[i]) for i in idx if freq[i] > 1}
    return vec, top


def aplicar_char_ngram(df: pd.DataFrame, col_txt, top_gramas: dict) -> np.ndarray:
    if not top_gramas or not col_txt: return np.zeros(len(df))
    cols = [c for c in col_txt if c in df.columns]
    if not cols: return np.zeros(len(df))
    blob = df[cols].fillna("").astype(str).agg(" ".join, axis=1)
    blob = blob.str.lower().str.replace(r"[^a-z0-9]", "", regex=True)
    top_set = set(top_gramas.keys())
    def _score(s: str) -> float:
        hits = sum(1 for g in top_set if g in s)
        return round(min(hits / 10.0, 1.0), 4)
    return blob.apply(_score).values


def calcular_outliers(df: pd.DataFrame) -> np.ndarray:
    num_cols = [c for c in df.columns
                if c.startswith("FE_") or c.startswith("KPI_")
                or c in ("ValorImportado", "PesoKg", "CapitalSocial")]
    num_cols = [c for c in num_cols if pd.api.types.is_numeric_dtype(df[c])]
    if not num_cols or len(df) < 10: return np.zeros(len(df))
    X = df[num_cols].fillna(0).values
    try:
        lof = LocalOutlierFactor(n_neighbors=min(20, len(df) - 1), contamination=0.1)
        pred = lof.fit_predict(X)
        s = -lof.negative_outlier_factor_
        s = (s - s.min()) / (s.max() - s.min() + 1e-9)
        return np.where(pred == -1, s, s * 0.3)
    except Exception:
        return np.zeros(len(df))


def calcular_risco(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["KPI_Outlier"] = calcular_outliers(df)
    mapa = {
        "subvalorizacao": "KPI_Subvalorizacao",
        "disparidade_hs": "KPI_DisparidadeHS",
        "surge_volume": "KPI_SurgeVolume",
        "desvio_eta": "KPI_DesvioETA",
        "completude": "KPI_Completude",
        "outlier": "KPI_Outlier",
        "consist_cadastral": "KPI_Ext_ConsistenciaCadastral",
        "idade_empresa": "KPI_Ext_IdadeEmpresa",
        "score_credito": "KPI_Ext_ScoreCredito",
        "padrao_char": "KPI_PadraoChar",
    }
    score = np.zeros(len(df))
    for k, w in PESO_KPIS.items():
        c = mapa.get(k)
        if c and c in df.columns:
            score += w * df[c].fillna(0).values
    df["RiskScore"] = np.clip(score * 100, 0, 100).round(2)
    bins = [0, 25, 50, 75, 100.01]
    labels = ["BAIXO", "MEDIO", "ALTO", "CRITICO"]
    cores = {"BAIXO": "#C6EFCE", "MEDIO": "#FFEB9C",
             "ALTO": "#FFC7CE", "CRITICO": "#C00000"}
    df["RiskLevel"] = pd.cut(df["RiskScore"], bins=bins, labels=labels, right=False)
    df["RiskColor"] = df["RiskLevel"].map(cores).fillna("#FFFFFF")
    return df


# =====================================================================
#  PREPROCESSADOR / ENSEMBLE
# =====================================================================
class TextConcat(BaseEstimator):
    def __init__(self, colunas): self.colunas = colunas
    def fit(self, X, y=None): return self
    def transform(self, X):
        if isinstance(X, pd.DataFrame):
            cols = [c for c in self.colunas if c in X.columns]
            if not cols: return np.array([""] * len(X))
            return X[cols].fillna("").astype(str).agg(" ".join, axis=1).values
        return np.array([" ".join(map(str, r)) for r in X])


class ColunasNumericas(BaseEstimator):
    def __init__(self, colunas): self.colunas = colunas
    def fit(self, X, y=None): return self
    def transform(self, X):
        if isinstance(X, pd.DataFrame):
            cols = [c for c in self.colunas if c in X.columns]
            if not cols: return np.zeros((len(X), 0))
            return X[cols].apply(pd.to_numeric, errors="coerce").fillna(0).values
        return np.asarray(X)


def construir_preprocessador(df_sample: pd.DataFrame, sel: dict) -> ColumnTransformer:
    txt = [c for c in sel.get("texto", []) if c in df_sample.columns]
    cat = [c for c in sel.get("categoricas", []) if c in df_sample.columns]
    num = [c for c in sel.get("numericas", []) if c in df_sample.columns]
    kpis = [c for c in COLUNAS_KPI if c in df_sample.columns]
    fe = [c for c in df_sample.columns if c.startswith("FE_")]
    num = list(dict.fromkeys(num + kpis + fe))

    transformers = []
    if txt:
        transformers.append(("hash_word", Pipeline([
            ("concat", TextConcat(txt)),
            ("hash", HashingVectorizer(n_features=HASH_WORD_FEATURES,
                                       ngram_range=(1, 2), alternate_sign=False,
                                       norm=None, lowercase=True)),
            ("norm", Normalizer(copy=False)),
        ]), txt))
        transformers.append(("hash_char", Pipeline([
            ("concat", TextConcat(txt)),
            ("hash", HashingVectorizer(n_features=HASH_CHAR_FEATURES,
                                       analyzer="char_wb", ngram_range=(3, 5),
                                       alternate_sign=False, norm=None,
                                       lowercase=True)),
            ("norm", Normalizer(copy=False)),
        ]), txt))
    if cat:
        transformers.append(("onehot", Pipeline([
            ("imp", SimpleImputer(strategy="constant", fill_value="NA")),
            ("oh", OneHotEncoder(handle_unknown="ignore", min_frequency=2)),
        ]), cat))
    if num:
        transformers.append(("num", Pipeline([
            ("sel", ColunasNumericas(num)),
            ("imp", SimpleImputer(strategy="median")),
            ("sc", StandardScaler()),
        ]), num))
    return ColumnTransformer(transformers, remainder="drop", sparse_threshold=0.3)


class NT1Ensemble(BaseEstimator, ClassifierMixin):
    def __init__(self, alpha: float = 1e-4, random_state: int = 42):
        self.alpha = alpha; self.random_state = random_state
        self._trained = {}; self._weights = {}

    def _parcial(self) -> dict:
        base = dict(alpha=self.alpha, max_iter=1, tol=None,
                    learning_rate="adaptive", eta0=0.01,
                    class_weight="balanced", average=True,
                    random_state=self.random_state)
        return {
            "sgd_log": SGDClassifier(loss="log_loss", **base),
            "sgd_huber": SGDClassifier(loss="modified_huber", **base),
            "pa": PassiveAggressiveClassifier(C=0.5, max_iter=1, tol=None,
                                              class_weight="balanced", average=True,
                                              random_state=self.random_state),
            "cnb": ComplementNB(alpha=0.3),
        }

    def _completo(self) -> dict:
        base = dict(alpha=self.alpha, max_iter=2000, tol=1e-4,
                    learning_rate="adaptive", eta0=0.01,
                    early_stopping=True, validation_fraction=0.1,
                    n_iter_no_change=10, class_weight="balanced",
                    average=True, random_state=self.random_state)
        return {
            "sgd_log": SGDClassifier(loss="log_loss", **base),
            "sgd_huber": SGDClassifier(loss="modified_huber", **base),
            "pa": PassiveAggressiveClassifier(C=1.0, max_iter=2000,
                                              class_weight="balanced", average=True,
                                              random_state=self.random_state),
            "cnb": ComplementNB(alpha=0.3),
            "svc": CalibratedClassifierCV(
                LinearSVC(C=1.0, class_weight="balanced", max_iter=5000,
                          random_state=self.random_state),
                cv=3, method="sigmoid"),
        }

    def partial_fit(self, X, y, classes=None):
        if classes is not None: self.classes_ = classes
        if not self._trained:
            self._trained = self._parcial()
            self._weights = {k: 1.0 for k in self._trained}
        for nome, m in self._trained.items():
            try:
                if not hasattr(m, "classes_"): m.partial_fit(X, y, classes=self.classes_)
                else: m.partial_fit(X, y)
            except Exception: pass
        try:
            for nome, m in self._trained.items():
                try:
                    p = m.predict(X)
                    acc = accuracy_score(y, p)
                    self._weights[nome] = 0.7 * self._weights.get(nome, 1.0) + 0.3 * acc
                except Exception: pass
            tot = sum(self._weights.values()) or 1.0
            self._weights = {k: v / tot for k, v in self._weights.items()}
        except Exception: pass
        return self

    def fit(self, X, y):
        self.classes_ = np.unique(y)
        self._trained = self._completo(); self._weights = {}
        n_min = np.min(np.bincount(y)) if len(self.classes_) > 1 else 0
        cv = min(3, n_min) if n_min >= 2 else None
        for nome, m in self._trained.items():
            try:
                m.fit(X, y); w = 1.0
                if cv and cv >= 2:
                    try:
                        sc = cross_val_score(clone(m), X, y, cv=cv,
                                             scoring="accuracy", n_jobs=-1)
                        w = float(np.clip(sc.mean(), 0.1, 1.0))
                    except Exception: pass
                self._weights[nome] = w
            except Exception as e:
                UI.warn(f"{nome}: {e}")
        self._trained = {k: v for k, v in self._trained.items() if k in self._weights}
        if not self._trained: raise RuntimeError("Nenhum modelo treinou.")
        tot = sum(self._weights.values()) or 1.0
        self._weights = {k: v / tot for k, v in self._weights.items()}
        return self

    def predict_proba(self, X, batch_size: int = 2000):
        n = X.shape[0]; n_c = len(self.classes_)
        probs = np.zeros((n, n_c))
        for i in range(0, n, batch_size):
            Xb = X[i:i + batch_size]
            pb = np.zeros((Xb.shape[0], n_c)); total_w = 0.0
            for nome, m in self._trained.items():
                if not hasattr(m, "predict_proba"): continue
                try: p = m.predict_proba(Xb)
                except Exception: continue
                if not np.array_equal(m.classes_, self.classes_):
                    full = np.zeros((Xb.shape[0], n_c))
                    for j, c in enumerate(m.classes_):
                        idx = np.where(self.classes_ == c)[0]
                        if len(idx): full[:, idx[0]] = p[:, j]
                    p = full
                w = self._weights.get(nome, 1.0)
                pb += p * w; total_w += w
            if total_w == 0: pb[:, 0] = 1.0
            else: pb /= pb.sum(axis=1, keepdims=True) + 1e-12
            probs[i:i + batch_size] = pb
        return probs

    def predict(self, X, batch_size: int = 2000):
        n = X.shape[0]
        preds = np.zeros(n, dtype=self.classes_.dtype)
        for i in range(0, n, batch_size):
            Xb = X[i:i + batch_size]
            pb = np.zeros((Xb.shape[0], len(self.classes_))); total_w = 0.0
            for nome, m in self._trained.items():
                if not hasattr(m, "predict_proba"): continue
                try: p = m.predict_proba(Xb)
                except Exception: continue
                if not np.array_equal(m.classes_, self.classes_):
                    full = np.zeros((Xb.shape[0], len(self.classes_)))
                    for j, c in enumerate(m.classes_):
                        idx = np.where(self.classes_ == c)[0]
                        if len(idx): full[:, idx[0]] = p[:, j]
                    p = full
                w = self._weights.get(nome, 1.0)
                pb += p * w; total_w += w
            if total_w == 0: pb[:, 0] = 1.0
            preds[i:i + batch_size] = self.classes_[np.argmax(pb, axis=1)]
        return preds


# =====================================================================
#  NEURAL TRAINER + ADAPTER
# =====================================================================
class NeuralTrainer:
    SCRIPT_BASE = [
        "import numpy as np",
        "from sklearn.neural_network import MLPClassifier",
        "from sklearn.decomposition import TruncatedSVD",
        "from sklearn.metrics import accuracy_score, f1_score",
        "",
        "def treinar_rede_neural(X, y, hidden, lr, epochs, seed):",
        "    # 1) reduzir dimensionalidade (esparso -> denso)",
        "    svd = TruncatedSVD(n_components=SVD_DIM, random_state=seed)",
        "    X_dense = svd.fit_transform(X)",
        "    ",
        "    # 2) split treino/validacao",
        "    X_tr, X_va, y_tr, y_va = train_test_split(X_dense, y, test_size=0.2)",
        "    ",
        "    # 3) montar a rede",
        "    mlp = MLPClassifier(hidden_layer_sizes=hidden,",
        "                        learning_rate_init=lr,",
        "                        batch_size=BATCH, max_iter=1,",
        "                        warm_start=True, random_state=seed)",
        "    ",
        "    # 4) loop de treino",
        "    for epoch in range(epochs):",
        "        mlp.fit(X_tr, y_tr)          # 1 epoca",
        "        loss = mlp.loss_",
        "        acc  = accuracy_score(y_tr, mlp.predict(X_tr))",
        "        vacc = accuracy_score(y_va, mlp.predict(X_va))",
        "        yield epoch, loss, acc, vacc",
    ]

    def __init__(self, cfg: Config, inv_classes: dict):
        self.cfg = cfg
        self.inv_classes = inv_classes
        self.viewer = CodeViewer(
            titulo="neural_train.py",
            max_vis=18,
            delay=cfg.neural_delay if cfg.neural_show_code else 0.0,
            enabled=cfg.neural_show_code)
        self.tabela = LiveTable(
            colunas=["Epoch", "Loss", "Acc Treino", "Acc Val", "Gap", "Tempo"],
            titulo="Treinamento Rede Neural  ·  MLPClassifier",
            larguras=[7, 12, 12, 12, 10, 9])
        self.modelo: Optional[MLPClassifier] = None
        self.svd: Optional[TruncatedSVD] = None
        self.metricas: dict = {}

    def _preparar_dados(self, X, y):
        max_rows = self.cfg.neural_max_rows
        if X.shape[0] > max_rows:
            idx = np.random.RandomState(RANDOM_STATE).choice(
                X.shape[0], max_rows, replace=False)
            X = X[idx]; y = y[idx]
            UI.info(f"Rede neural: subsample para {max_rows} linhas")
        svd_dim = self.cfg.neural_svd
        if svd_dim > 0 and X.shape[1] > svd_dim:
            self.viewer.escrever(
                f"svd = TruncatedSVD(n_components={svd_dim}, random_state={RANDOM_STATE})", delay=0)
            self.viewer.escrever("X_dense = svd.fit_transform(X)", delay=0)
            self.svd = TruncatedSVD(n_components=svd_dim, random_state=RANDOM_STATE)
            Xd = self.svd.fit_transform(X)
            self.viewer.escrever(
                f"# X_dense.shape = {Xd.shape}  (var_expl={self.svd.explained_variance_ratio_.sum():.4f})",
                delay=0)
        else:
            try: Xd = X.toarray() if hasattr(X, "toarray") else np.asarray(X)
            except Exception: Xd = np.asarray(X)
        return Xd, y

    def treinar(self, X, y) -> dict:
        UI.secao("F3.5", "Rede Neural  ·  Treino Supervisionado")
        log_jsonl("neural_inicio",
                  n_amostras=int(X.shape[0]), n_features=int(X.shape[1]),
                  hidden=str(self.cfg.hidden_tuple()),
                  lr=self.cfg.neural_lr, epochs=self.cfg.neural_epochs)

        self.viewer.abrir()
        for l in self.SCRIPT_BASE[:4]:
            self.viewer.escrever(l, delay=0.0)
        self.viewer.pausa(0.05)
        self.viewer.bloco(self.SCRIPT_BASE[4:9])
        self.viewer.pausa(0.1)

        Xd, y = self._preparar_dados(X, y)
        if len(np.unique(y)) < 2:
            self.viewer.escrever("# nenhuma classe valida encontrada", delay=0)
            self.viewer.fim("abortado")
            return {}
        self.viewer.bloco(self.SCRIPT_BASE[9:15], delay=0.0)
        self.viewer.pausa(0.05)
        self.viewer.bloco(self.SCRIPT_BASE[15:19], delay=0.0)

        try:
            strat = y if np.min(np.bincount(y)) >= 2 else None
            X_tr, X_va, y_tr, y_va = train_test_split(
                Xd, y, test_size=0.2, random_state=RANDOM_STATE, stratify=strat)
        except Exception:
            X_tr, X_va, y_tr, y_va = train_test_split(
                Xd, y, test_size=0.2, random_state=RANDOM_STATE)
        self.viewer.escrever(f"# X_tr={X_tr.shape}  X_va={X_va.shape}", delay=0)

        hidden = self.cfg.hidden_tuple()
        mlp = MLPClassifier(
            hidden_layer_sizes=hidden,
            learning_rate_init=self.cfg.neural_lr,
            batch_size=min(NEURAL_BATCH, max(16, len(X_tr) // 20)),
            max_iter=1, warm_start=True,
            early_stopping=False,
            random_state=RANDOM_STATE,
            activation="relu", solver="adam",
        )
        self.viewer.escrever(f"mlp = MLPClassifier(hidden_layer_sizes={hidden})", delay=0)
        self.viewer.escrever(f"mlp.learning_rate_init = {self.cfg.neural_lr}", delay=0)

        self.viewer.bloco([
            "",
            f"for epoch in range({self.cfg.neural_epochs}):",
            "    mlp.fit(X_tr, y_tr)",
            "    loss = mlp.loss_",
            "    acc  = accuracy_score(y_tr, mlp.predict(X_tr))",
            "    vacc = accuracy_score(y_va, mlp.predict(X_va))",
            "    yield epoch, loss, acc, vacc",
            "",
        ], delay=0.0)

        self.tabela.abrir()
        best_vacc = 0.0; best_loss = float("inf")
        hist = []

        for ep in range(self.cfg.neural_epochs):
            t0 = time.time()
            try:
                mlp.fit(X_tr, y_tr)
                loss = float(getattr(mlp, "loss_", np.nan))
                acc  = float(accuracy_score(y_tr, mlp.predict(X_tr)))
                vacc = float(accuracy_score(y_va, mlp.predict(X_va)))
            except Exception as e:
                self.viewer.escrever(f"# ERRO epoch {ep}: {e}", delay=0)
                log_jsonl("neural_epoch_erro", level="WARN", epoch=ep, erro=str(e))
                continue
            dt = time.time() - t0
            gap = acc - vacc

            self.viewer.escrever(
                f"epoch={ep:>3}  loss={loss:.5f}  acc={acc:.4f}  vacc={vacc:.4f}",
                delay=0)

            self.tabela.add([ep + 1, loss, acc, vacc, gap, f"{dt:.2f}s"])
            hist.append({"epoch": ep + 1, "loss": loss, "acc": acc,
                         "vacc": vacc, "gap": gap, "tempo_s": round(dt, 3)})
            log_jsonl("neural_epoch", **hist[-1])

            if vacc > best_vacc:
                best_vacc = vacc
                best_loss = loss

        self.modelo = mlp
        self.metricas = {
            "epochs": len(hist),
            "best_val_acc": round(best_vacc, 4),
            "best_loss": round(best_loss, 4),
            "final_train_acc": round(hist[-1]["acc"], 4) if hist else 0.0,
            "final_val_acc": round(hist[-1]["vacc"], 4) if hist else 0.0,
            "historico": hist[-25:],
        }

        self.viewer.bloco([
            "",
            "# treino concluido",
            f"# best_val_acc = {best_vacc:.4f}",
            f"# best_loss    = {best_loss:.5f}",
        ], delay=0.0)
        self.viewer.fim(f"MLP treinada em {len(hist)} epochs")

        self.tabela.fim(
            f"Rede neural concluida  ·  best val_acc = {best_vacc:.4f}  "
            f"loss = {best_loss:.5f}")
        UI.fim_secao()
        log_jsonl("neural_fim", **{k: v for k, v in self.metricas.items()
                                    if k != "historico"})
        return self.metricas

    def predict_proba(self, X) -> np.ndarray:
        if self.modelo is None: return None
        Xd = self.svd.transform(X) if self.svd is not None else \
            (X.toarray() if hasattr(X, "toarray") else np.asarray(X))
        return self.modelo.predict_proba(Xd)

    def predict(self, X) -> np.ndarray:
        if self.modelo is None: return None
        Xd = self.svd.transform(X) if self.svd is not None else \
            (X.toarray() if hasattr(X, "toarray") else np.asarray(X))
        return self.modelo.predict(Xd)


class NeuralAdapter(BaseEstimator, ClassifierMixin):
    def __init__(self, trainer: NeuralTrainer, classes_: np.ndarray):
        self.trainer = trainer
        self.classes_ = np.asarray(classes_)
    def predict_proba(self, X):
        return self.trainer.predict_proba(X)
    def predict(self, X):
        return self.trainer.predict(X)


# =====================================================================
#  TREINADOR CONTINUO (com stream enriquecido)
# =====================================================================
class ParadaControlada:
    def __init__(self):
        self.event = threading.Event()
        threading.Thread(target=self._loop, daemon=True).start()
    def _loop(self):
        while not self.event.is_set():
            try: line = input()
            except (EOFError, Exception): break
            if line.strip().lower() in ("parar", "stop", "s", "q", "quit", "exit"):
                self.event.set(); UI.warn("Parada solicitada..."); break
    def parado(self) -> bool: return self.event.is_set()


class TreinadorContinuo:
    def __init__(self, modelo, preproc, classes_map, pasta_dados, col_id,
                 col_txt, col_cat, top_gramas, modo_completo,
                 stream_enriquecido: bool = False,
                 stream_budget: int = 500,
                 stream_so_cache: bool = False):
        self.modelo = modelo; self.preproc = preproc
        self.classes_map = classes_map
        self.inv_classes = {v: k for k, v in classes_map.items()}
        self.pasta_dados = pasta_dados; self.col_id = col_id
        self.col_txt = col_txt; self.col_cat = col_cat
        self.top_gramas = top_gramas; self.modo_completo = modo_completo
        self.parada = ParadaControlada()
        self.historico = []; self.epoch = 0
        self.amostrador = Amostrador(SAMPLE_ROWS)
        self.enriquecidos = set(); self.metricas_ultima = {}

        # stream enriquecido
        self.stream_enriquecido = stream_enriquecido
        self.stream_budget      = stream_budget
        self.stream_so_cache    = stream_so_cache
        self._budget_usado      = 0
        self._budget_lock       = threading.Lock()
        self._stats_enriq       = {"chunks_com_ext": 0, "chunks_sem_ext": 0,
                                   "cnpjs_novos": 0, "api_calls": 0}

    def _preparar_chunk(self, chunk):
        if self.col_id not in chunk.columns: return None
        chunk = chunk.copy()
        chunk[self.col_id] = normalizar_id(chunk[self.col_id])
        chunk = chunk[chunk[self.col_id].astype(str).str.len() > 0]
        if len(chunk) == 0: return None

        if self.stream_enriquecido and CAMPO_ID_DEFAULT in chunk.columns:
            cnpjs_novos = set(chunk[CAMPO_ID_DEFAULT].astype(str).unique()) - self.enriquecidos
            if cnpjs_novos:
                with self._budget_lock:
                    pode_api = (not self.stream_so_cache) and \
                               (self._budget_usado < self.stream_budget)
                    budget_restante = self.stream_budget - self._budget_usado
                so_cache = self.stream_so_cache or not pode_api
                cota = 0 if so_cache else min(budget_restante, MAX_CNPJ_EXTERNO)

                chunk = enriquecer_cnpjs(chunk, max_cnpjs=cota, so_cache=so_cache)
                self.enriquecidos.update(cnpjs_novos)

                with self._budget_lock:
                    self._budget_usado += cota
                    self._stats_enriq["cnpjs_novos"] += len(cnpjs_novos)
                    self._stats_enriq["api_calls"]   += cota
                    if "EXT_Situacao" in chunk.columns and chunk["EXT_Situacao"].notna().any():
                        self._stats_enriq["chunks_com_ext"] += 1
                    else:
                        self._stats_enriq["chunks_sem_ext"] += 1
            else:
                chunk = enriquecer_cnpjs(chunk, max_cnpjs=0, so_cache=True)

        elif self.modo_completo and CAMPO_ID_DEFAULT in chunk.columns:
            novos = set(chunk[CAMPO_ID_DEFAULT].unique()) - self.enriquecidos
            if novos:
                chunk = enriquecer_cnpjs(chunk)
                self.enriquecidos.update(novos)

        chunk = adicionar_features_avancadas(chunk, self.col_txt, self.col_cat)
        chunk = calcular_kpis_negocio(chunk)
        chunk["KPI_PadraoChar"] = aplicar_char_ngram(chunk, self.col_txt, self.top_gramas)
        chunk = calcular_risco(chunk)
        return chunk

    def _filtrar_conhecidos(self, chunk):
        ids = chunk[self.col_id].astype(str).values
        mask = np.array([i in self.classes_map for i in ids])
        return chunk[mask].copy(), ids[mask]

    def treinar(self, max_epochs=0):
        UI.secao("F2", "Treinamento Continuo")
        UI.info("Digite 'parar' + Enter para encerrar a qualquer momento.")
        if max_epochs: UI.info(f"Limite: {max_epochs} epoch(s).")
        n_classes = len(self.classes_map)
        classes_arr = np.arange(n_classes)
        log_jsonl("treino_inicio", n_classes=n_classes, max_epochs=max_epochs)

        while not self.parada.parado():
            self.epoch += 1
            accs, f1s = [], []
            n_chunks = 0; n_linhas = 0; t_ep = time.time()
            log_jsonl("epoch_inicio", epoch=self.epoch)
            for _, raw in iterar_chunks(self.pasta_dados):
                if self.parada.parado(): break
                try:
                    chunk = self._preparar_chunk(raw)
                    if chunk is None or len(chunk) == 0: continue
                    chunk, _ = self._filtrar_conhecidos(chunk)
                    if len(chunk) == 0: continue
                    y = np.array([self.classes_map[c]
                                  for c in chunk[self.col_id].astype(str).values])
                    X = self.preproc.transform(chunk)
                    self.modelo.partial_fit(X, y, classes=classes_arr)
                    self.amostrador.add(chunk.head(min(len(chunk), 200)))
                    try:
                        pred = self.modelo.predict(X)
                        a = accuracy_score(y, pred)
                        f = f1_score(y, pred, average="macro", zero_division=0)
                        accs.append(a); f1s.append(f)
                        log_jsonl("chunk", epoch=self.epoch, n=len(chunk),
                                  acc=round(float(a), 4), f1=round(float(f), 4))
                    except Exception: pass
                    n_chunks += 1; n_linhas += len(chunk)
                    UI.progresso(n_chunks, max(n_chunks, n_chunks + 1),
                                 f"Epoch {self.epoch}",
                                 extra=f"{n_linhas} linhas")
                except Exception as e:
                    UI.warn(f"chunk: {e}")
                    log_jsonl("chunk_erro", level="WARN", erro=str(e))
                    continue
            UI.fim_progresso()
            if not accs:
                UI.warn(f"Epoch {self.epoch}: sem dados validos.")
                if max_epochs and self.epoch >= max_epochs: break
                continue
            dt = time.time() - t_ep
            self.metricas_ultima = {
                "epoch": self.epoch,
                "accuracy": round(float(np.mean(accs)), 4),
                "f1_macro": round(float(np.mean(f1s)), 4),
                "chunks": n_chunks, "linhas": n_linhas,
                "tempo_s": round(dt, 2),
            }
            self.historico.append(self.metricas_ultima)
            UI.tabela(
                ["Epoch", "Accuracy", "F1 macro", "Chunks", "Linhas", "Tempo"],
                [[self.epoch, f"{np.mean(accs):.4f}", f"{np.mean(f1s):.4f}",
                  n_chunks, n_linhas, f"{dt:.1f}s"]])
            log_jsonl("epoch_fim", **self.metricas_ultima)
            if max_epochs and self.epoch >= max_epochs: break

        UI.fim_secao()
        if self.stream_enriquecido:
            UI.caixa("Stream enriquecido  ·  resumo",
                     f"Chunks com EXT_*:   {self._stats_enriq['chunks_com_ext']}\n"
                     f"Chunks sem EXT_*:   {self._stats_enriq['chunks_sem_ext']}\n"
                     f"CNPJs unicos novos: {self._stats_enriq['cnpjs_novos']}\n"
                     f"Chamadas API:       {self._stats_enriq['api_calls']} "
                     f"(budget {self.stream_budget})",
                     cor=UI.GREEN)
            log_jsonl("stream_enriquecido_fim", **self._stats_enriq,
                      budget=self.stream_budget, budget_usado=self._budget_usado)
        log_jsonl("treino_fim", epochs_executadas=len(self.historico))
        return {"historico": self.historico, "ultima": self.metricas_ultima}


# =====================================================================
#  AVALIACAO / HISTORICO / PADROES
# =====================================================================
def avaliar_modelo(modelo, X, y) -> dict:
    if len(y) == 0: return {}
    metricas = {}
    try:
        y_pred = modelo.predict(X)
        metricas["accuracy"]  = round(float(accuracy_score(y, y_pred)), 4)
        metricas["f1_macro"]  = round(float(f1_score(y, y_pred, average="macro", zero_division=0)), 4)
        metricas["precision"] = round(float(precision_score(y, y_pred, average="macro", zero_division=0)), 4)
        metricas["recall"]    = round(float(recall_score(y, y_pred, average="macro", zero_division=0)), 4)
    except Exception:
        return metricas
    try:
        n_classes = len(modelo.classes_)
        labels = modelo.classes_
        y_proba = modelo.predict_proba(X)
        for k in TOP_K:
            if k >= n_classes: continue
            try:
                metricas[f"top_{k}"] = round(float(
                    top_k_accuracy_score(y, y_proba, k=k, labels=labels)), 4)
            except Exception: pass
        try:
            metricas["log_loss"] = round(float(
                log_loss(y, y_proba, labels=labels)), 4)
        except Exception: pass
    except Exception: pass
    return metricas


def carregar_historico() -> list:
    os.makedirs(PASTA_HIST, exist_ok=True)
    if os.path.exists(HIST_PATH):
        try:
            with open(HIST_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception: return []
    return []


def salvar_sessao(sessao: dict) -> list:
    hist = carregar_historico(); hist.append(sessao)
    with open(HIST_PATH, "w", encoding="utf-8") as f:
        json.dump(hist, f, indent=2, ensure_ascii=False, default=str)
    return hist


def detectar_padrao(hist: list) -> dict:
    if len(hist) < 3:
        return {"aviso": f"Historico curto ({len(hist)} sessoes)."}
    ult = hist[-min(len(hist), N_SESSOES_PADRAO):]
    accs = [s.get("metricas", {}).get("accuracy", 0) for s in ult]
    tops = [s.get("metricas", {}).get("top_5", 0) for s in ult]
    tendencia = "estavel"
    if len(accs) >= 3:
        if accs[-1] > accs[0] + 0.02: tendencia = "melhora"
        elif accs[-1] < accs[0] - 0.02: tendencia = "degradacao"
    kpi_counter = Counter()
    for s in ult:
        for k, v in (s.get("kpis_medio_top10") or {}).items():
            if v > 0.5: kpi_counter[k] += 1
    limiar = 0.6 * len(ult)
    kpis_rec = [k for k, c in kpi_counter.items() if c >= limiar]
    return {
        "sessoes_analisadas": len(ult),
        "accuracy_media": round(float(np.mean(accs)), 4),
        "accuracy_std": round(float(np.std(accs)), 4),
        "top5_media": round(float(np.mean(tops)), 4),
        "tendencia": tendencia,
        "kpis_recorrentes": kpis_rec,
        "assinatura": hashlib.md5(",".join(sorted(kpis_rec)).encode()).hexdigest()[:12],
        "pronto_para_palpite": len(hist) >= N_SESSOES_PADRAO,
    }


def salvar_padrao(p: dict):
    os.makedirs(PASTA_HIST, exist_ok=True)
    with open(PADRAO_PATH, "w", encoding="utf-8") as f:
        json.dump(p, f, indent=2, ensure_ascii=False, default=str)


def adicionar_resultados(df, modelo, preproc, inv_classes, col_id,
                         col_txt, col_cat, top_gramas, batch_size=2000) -> pd.DataFrame:
    df = df.copy()
    df[col_id] = normalizar_id(df[col_id])
    df = adicionar_features_avancadas(df, col_txt, col_cat)
    df = calcular_kpis_negocio(df)
    df["KPI_PadraoChar"] = aplicar_char_ngram(df, col_txt, top_gramas)
    df = calcular_risco(df)
    X = preproc.transform(df)
    n = X.shape[0]; n_c = len(modelo.classes_)
    proba = np.zeros((n, n_c))
    for i in range(0, n, batch_size):
        Xb = X[i:i + batch_size]
        try: pb = modelo.predict_proba(Xb)
        except Exception:
            pb = np.zeros((Xb.shape[0], n_c)); pb[:, 0] = 1.0
        proba[i:i + batch_size] = pb
    idx_pred = proba.argmax(axis=1)
    df["RES_Previsto"]  = [inv_classes.get(int(i), str(i)) for i in idx_pred]
    df["RES_Confianca"] = np.round(proba.max(axis=1), 4)
    top_k = min(5, n_c)
    top_idx = np.argsort(proba, axis=1)[:, ::-1][:, :top_k]
    for r in range(top_k):
        df[f"RES_Top{r+1}"] = [inv_classes.get(int(i), str(i)) for i in top_idx[:, r]]
        df[f"RES_Top{r+1}_p"] = np.round(
            np.take_along_axis(proba, top_idx[:, r:r+1], axis=1).ravel(), 4)
    df["RES_Entropia"] = np.round(-np.sum(proba * np.log(proba + 1e-12), axis=1), 4)
    sp = np.sort(proba, axis=1)[:, ::-1]
    df["RES_Margem"] = np.round(sp[:, 0] - sp[:, 1], 4) if sp.shape[1] >= 2 else 0.0
    return df


def processar_e_escrever(pasta_dados, modelo, preproc, inv_classes, col_id,
                         col_txt, col_cat, top_gramas, modo_completo) -> list:
    os.makedirs(PASTA_PARTS, exist_ok=True)
    for f in Path(PASTA_PARTS).glob("*.parquet"):
        try: f.unlink()
        except Exception: pass
    enriquecidos = set(); parts = []
    for i, (nome, raw) in enumerate(iterar_chunks(pasta_dados)):
        if col_id not in raw.columns: continue
        try:
            chunk = raw.copy()
            chunk[col_id] = normalizar_id(chunk[col_id])
            chunk = chunk[chunk[col_id].astype(str).str.len() > 0]
            if len(chunk) == 0: continue
            if modo_completo and CAMPO_ID_DEFAULT in chunk.columns:
                novos = set(chunk[CAMPO_ID_DEFAULT].unique()) - enriquecidos
                if novos:
                    chunk = enriquecer_cnpjs(chunk); enriquecidos.update(novos)
            chunk = adicionar_resultados(chunk, modelo, preproc, inv_classes,
                                         col_id, col_txt, col_cat, top_gramas)
            part = os.path.join(PASTA_PARTS, f"part_{i:05d}.parquet")
            chunk.to_parquet(part, index=False, engine="pyarrow")
            parts.append(part)
            UI.progresso(i + 1, i + 2, "Gerando", extra=f"{len(chunk)} linhas")
        except Exception as e:
            UI.warn(f"parte {i}: {e}")
    UI.fim_progresso(f"{len(parts)} partes geradas")
    return parts


def exportar_xlsx(df, caminho):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils.dataframe import dataframe_to_rows
        wb = Workbook(); ws = wb.active; ws.title = "ExtractV3"
        cores = {"BAIXO": "C6EFCE", "MEDIO": "FFEB9C",
                 "ALTO": "FFC7CE", "CRITICO": "C00000"}
        cols = list(df.columns); ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        col_risk = cols.index("RiskLevel") + 1 if "RiskLevel" in cols else None
        for r in dataframe_to_rows(df, index=False, header=False):
            ws.append(r); ri = ws.max_row
            if col_risk:
                v = str(ws.cell(row=ri, column=col_risk).value or "").upper()
                fill = PatternFill("solid", fgColor=cores.get(v, "FFFFFF"))
                for c in range(1, len(cols) + 1):
                    ws.cell(row=ri, column=c).fill = fill
        wb.save(caminho); UI.ok(f"XLSX: {caminho}")
    except Exception as e:
        UI.warn(f"XLSX: {e}")


def exportar_top_risco(parts, caminho, n=N_TOP_RISCO_XLSX):
    dfs = []
    for p in parts:
        try:
            d = pd.read_parquet(p, engine="pyarrow")
            if "RiskScore" in d.columns: dfs.append(d)
        except Exception: continue
    if not dfs:
        UI.warn("Sem partes para top risco."); return
    df = pd.concat(dfs, ignore_index=True).sort_values(
        "RiskScore", ascending=False).head(n)
    exportar_xlsx(df, caminho)


# =====================================================================
#  CONSOLE INTERATIVO
# =====================================================================
def console_loop(modelo, preproc, inv_classes, col_txt, col_cat, top_gramas, col_id):
    while True:
        op = UI.menu("O que deseja agora?", [
            ("c", "Consultar  —  inferir para um registro digitado"),
            ("p", "Capturar PDF / clipboard"),
            ("h", "Ver padrao / historico"),
            ("m", "Metricas do modelo atual"),
            ("s", "Sair"),
        ])
        if op in ("s", "", "0"): break
        if op == "c":
            campos = col_txt + col_cat
            valores = coletar_valores_alvo(campos)
            if valores is None: continue
            try:
                base = pd.DataFrame([valores])
                base[col_id] = "00000000000000"
                base = adicionar_features_avancadas(base, col_txt, col_cat)
                base = calcular_kpis_negocio(base)
                base["KPI_PadraoChar"] = aplicar_char_ngram(base, col_txt, top_gramas)
                base = calcular_risco(base)
                Xq = preproc.transform(base)
                probs = modelo.predict_proba(Xq)[0]
                top = np.argsort(probs)[::-1][:5]
                linhas = "\n".join(
                    f"{r+1}o  {inv_classes.get(int(i), i)}  (p={probs[i]:.4f})"
                    for r, i in enumerate(top))
                alertar(f"Top-5 previstos:\n\n{linhas}", "Consulta")
            except Exception as e:
                alertar(f"Erro: {e}", "Aviso")
        elif op == "p":
            txt = captura_clipboard()
            if not txt: continue
            raw = re.findall(r"\b\d{14}\b", re.sub(r"\D", "", txt))
            fmt = re.findall(r"\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}", txt)
            todos = list(dict.fromkeys(
                [c.zfill(14) for c in raw] +
                [re.sub(r"\D", "", c).zfill(14) for c in fmt]))
            linhas = "\n".join(f"  - {c}" for c in todos[:20])
            alertar(f"Captura\n\nCNPJs: {len(todos)}\n{linhas}", "PDF")
        elif op == "h":
            hist = carregar_historico()
            if not hist:
                alertar("Sem historico.", "Historico"); continue
            linhas_t = []
            for i, s in enumerate(hist[-10:], 1):
                m = s.get("metricas", {})
                linhas_t.append([f"#{i}", s.get("timestamp", "-"),
                                 m.get("accuracy", "-"), m.get("top_5", "-")])
            UI.tabela(["Sessao", "Timestamp", "Accuracy", "Top-5"], linhas_t)
            pad = detectar_padrao(hist)
            UI.caixa("Padrao",
                     f"Tendencia: {pad.get('tendencia','-')}\n"
                     f"KPIs recorrentes: {pad.get('kpis_recorrentes', [])}\n"
                     f"Pronto p/ palpite: {pad.get('pronto_para_palpite', False)}")
        elif op == "m":
            UI.caixa("Metricas do modelo atual",
                     json.dumps(getattr(modelo, "_weights", {}),
                                indent=2, ensure_ascii=False))


# =====================================================================
#  MAIN
# =====================================================================
def _preparar_pastas(cfg: Config):
    for p in (cfg.pasta_saida, cfg.pasta_cache, cfg.pasta_hist,
              cfg.pasta_parts, cfg.pasta_logs):
        os.makedirs(p, exist_ok=True)


def _selecionar_colunas(cfg: Config, amostra: pd.DataFrame, col_id: str) -> dict:
    if cfg.headless and (cfg.colunas_texto or cfg.colunas_cat or cfg.colunas_num):
        return {
            "texto":      [c for c in cfg.colunas_texto if c in amostra.columns],
            "categoricas":[c for c in cfg.colunas_cat   if c in amostra.columns],
            "numericas":  [c for c in cfg.colunas_num   if c in amostra.columns],
            "temporais":  [], "ignorar": [],
        }
    if cfg.headless:
        cls = classificar_colunas(amostra, col_id)
        UI.warn("Headless sem --texto/--cat/--num: usando classificacao automatica.")
        return {
            "texto": cls.get("texto", []),
            "categoricas": cls.get("categoricas", []),
            "numericas": cls.get("numericas", []),
            "temporais": cls.get("temporais", []),
            "ignorar": cls.get("ignorar", []),
        }
    return selecionar_colunas_terminal(classificar_colunas(amostra, col_id), amostra)


def _perguntar_epochs(cfg: Config) -> int:
    if cfg.headless: return cfg.epochs
    UI.secao("F2", "Treinamento")
    r = UI.prompt("Quantas epochs? (0 = infinito)", "0")
    try: return int(r) if r else 0
    except Exception: return 0


def main(argv=None):
    cfg = parse_args(argv)
    UI.init()
    if cfg.no_color: UI._CORES_ATIVAS = False
    UI.init_jsonl(cfg.pasta_logs, habilitado=cfg.log_jsonl)
    _preparar_pastas(cfg)

    UI.banner(f"{APP_NOME}  ·  v{APP_VERSAO}", APP_SUB)
    if cfg.headless:
        UI.info("Modo HEADLESS ativo (sem interacao).")
        UI.info(f"modo={cfg.modo} epochs={cfg.epochs} "
                f"neural={'on' if cfg.neural else 'off'} "
                f"refinamento={'off' if cfg.sem_refinamento else 'on'} "
                f"stream={'on' if cfg.stream_enriquecido else 'off'}")
    log_jsonl("sessao_inicio", **cfg.to_dict())
    t0 = time.time()

    if cfg.headless:
        modo_completo = (cfg.modo == "completo")
    else:
        resp = menu_inicial()
        if resp in (None, "Cancelar"):
            UI.warn("Cancelado."); log_jsonl("sessao_cancelada"); return
        if resp == "Config":
            UI.caixa("Config",
                     f"Dados: {cfg.pasta_dados}\nSaida: {cfg.pasta_saida}\n"
                     f"Cache: {cfg.pasta_cache}\nHistorico: {cfg.pasta_hist}\n"
                     f"Rede Neural: {'on' if cfg.neural else 'off'} "
                     f"({cfg.neural_epochs} epochs)")
            resp = menu_inicial()
            if resp in (None, "Cancelar", "Config"): return
        modo_completo = (resp == "Completo")

    arquivos = listar_arquivos(cfg.pasta_dados)
    if not arquivos:
        UI.erro(f"Nenhum arquivo em {cfg.pasta_dados}")
        log_jsonl("erro", level="ERRO", motivo="sem_arquivos")
        if not cfg.headless: alertar(f"Nenhum arquivo em {cfg.pasta_dados}", "Aviso")
        else: sys.exit(1)
        return
    UI.secao("F0", "Amostragem e Deteccao")
    UI.info(f"{len(arquivos)} arquivo(s) encontrados em {cfg.pasta_dados}")

    with UI.Spinner("Carregando amostra..."):
        amostra = amostrar_dados(cfg.pasta_dados, n=cfg.sample_rows)
    if amostra.empty:
        alertar("Amostra vazia.", "Erro")
        log_jsonl("erro", level="ERRO", motivo="amostra_vazia")
        if cfg.headless: sys.exit(1)
        return

    col_id = detectar_coluna_id(amostra)
    if col_id is None:
        alertar("Nenhuma coluna ID.", "Erro")
        log_jsonl("erro", level="ERRO", motivo="sem_coluna_id")
        if cfg.headless: sys.exit(1)
        return
    UI.ok(f"Coluna ID: '{col_id}'")

    digitos = amostra[col_id].astype(str).str.replace(r"\D", "", regex=True)
    cnpj_like = 13 <= digitos.str.len().mean() <= 15
    amostra[col_id] = (normalizar_id_cnpj(amostra[col_id]) if cnpj_like
                       else normalizar_id(amostra[col_id]))
    amostra = amostra[amostra[col_id].astype(str).str.len() > 0]
    if amostra.empty:
        alertar("Amostra vazia apos limpeza.", "Erro")
        if cfg.headless: sys.exit(1)
        return
    UI.ok(f"Amostra: {len(amostra)} linhas × {amostra.shape[1]} colunas")
    UI.fim_secao()
    log_jsonl("amostra_ok", n_linhas=len(amostra),
              n_colunas=amostra.shape[1], col_id=col_id, cnpj_like=bool(cnpj_like))

    if cfg.dry_run:
        UI.caixa("DRY-RUN",
                 f"Coluna ID: {col_id}\nLinhas: {len(amostra)}\n"
                 f"Colunas: {amostra.shape[1]}\n"
                 f"Modo: {'completo' if modo_completo else 'rapido'}\n"
                 f"Epochs: {cfg.epochs}\n"
                 f"Neural: {'on' if cfg.neural else 'off'} "
                 f"({cfg.neural_epochs} epochs, hidden={cfg.hidden_tuple()})\n"
                 f"Stream enriquecido: {'on' if cfg.stream_enriquecido else 'off'}"
                 f"{' (so cache)' if cfg.stream_so_cache else ''} "
                 f"budget={cfg.stream_budget}\n"
                 f"Refinamento: {'off' if cfg.sem_refinamento else 'on'}\n"
                 f"XLSX: {'off' if cfg.sem_xlsx else 'on'}\n"
                 f"JSONL: {UI._jsonl_path}")
        log_jsonl("dry_run_ok", col_id=col_id, n_linhas=len(amostra))
        UI.banner("FIM (dry-run)"); return

    selecionadas = _selecionar_colunas(cfg, amostra, col_id)
    col_txt = selecionadas.get("texto", [])
    col_cat = selecionadas.get("categoricas", [])
    col_num = selecionadas.get("numericas", [])
    if not (col_txt or col_cat or col_num):
        alertar("Nenhuma coluna selecionada.", "Aviso")
        log_jsonl("erro", level="ERRO", motivo="sem_colunas")
        if cfg.headless: sys.exit(1)
        return
    log_jsonl("colunas_selecionadas",
              texto=len(col_txt), cat=len(col_cat), num=len(col_num))

    # F1
    UI.secao("F1", "Enriquecimento e Features")
    if modo_completo and CAMPO_ID_DEFAULT in amostra.columns:
        amostra = enriquecer_cnpjs(amostra, max_cnpjs=min(cfg.max_cnpj, len(amostra)))
    with UI.Spinner("Gerando features avancadas..."):
        amostra = adicionar_features_avancadas(amostra, col_txt, col_cat)
        amostra = calcular_kpis_negocio(amostra)
    with UI.Spinner("Ajustando char-ngram..."):
        _, top_gramas = fit_char_ngram(amostra, col_txt)
        amostra["KPI_PadraoChar"] = aplicar_char_ngram(amostra, col_txt, top_gramas)
    UI.ok(f"{len(top_gramas)} gramas recorrentes")
    log_jsonl("char_ngram", n_gramas=len(top_gramas))
    with UI.Spinner("Calculando risco..."):
        amostra = calcular_risco(amostra)
    with UI.Spinner("Ajustando preprocessador..."):
        preproc = construir_preprocessador(amostra, selecionadas)
        X_sample = preproc.fit_transform(amostra)
    UI.ok(f"Dimensao das features: {X_sample.shape}")
    log_jsonl("preprocessador", shape=list(X_sample.shape))

    ids_unicos = amostra[col_id].astype(str).unique()
    if len(ids_unicos) > cfg.max_classes:
        contagem = amostra[col_id].astype(str).value_counts()
        ids_unicos = contagem.head(cfg.max_classes).index.tolist()
        UI.warn(f"Limitando a {cfg.max_classes} classes")
    classes_map = {c: i for i, c in enumerate(ids_unicos)}
    inv_classes = {v: k for k, v in classes_map.items()}

    if len(classes_map) < 2:
        alertar("Menos de 2 classes.", "Aviso")
        log_jsonl("erro", level="ERRO", motivo="poucas_classes",
                  n_classes=len(classes_map))
        if cfg.headless: sys.exit(1)
        return
    UI.ok(f"{len(classes_map)} classes mapeadas")
    log_jsonl("classes_ok", n_classes=len(classes_map))

    y_sample = np.array([classes_map[c] for c in amostra[col_id].astype(str).values
                         if c in classes_map])
    mask_sample = np.array([c in classes_map for c in amostra[col_id].astype(str).values])
    X_fit = X_sample[mask_sample]

    with UI.Spinner("Inicializando ensemble..."):
        modelo = NT1Ensemble(alpha=1e-4, random_state=RANDOM_STATE)
        modelo.partial_fit(X_fit, y_sample, classes=np.arange(len(classes_map)))
    UI.ok("Ensemble inicializado (4 modelos parciais)")
    UI.fim_secao()

    # F2
    _cache_quente = os.path.exists(os.path.join(cfg.pasta_cache, "cnpj_cache.parquet"))
    _enriquecer_no_treino = modo_completo and _cache_quente
    _stream = cfg.stream_enriquecido and modo_completo

    if _stream:
        UI.info(f"Stream enriquecido ATIVO — budget={cfg.stream_budget} "
                f"chamadas, so_cache={cfg.stream_so_cache}")
    elif _enriquecer_no_treino:
        UI.info("Cache de CNPJs quente — enriquecimento via cache no treino.")
    else:
        UI.info("Treino sem enriquecimento (F2.5 enriquecera o reservoir).")

    treinador = TreinadorContinuo(
        modelo=modelo, preproc=preproc, classes_map=classes_map,
        pasta_dados=cfg.pasta_dados, col_id=col_id, col_txt=col_txt,
        col_cat=col_cat, top_gramas=top_gramas,
        modo_completo=_enriquecer_no_treino or _stream,
        stream_enriquecido=_stream,
        stream_budget=cfg.stream_budget,
        stream_so_cache=cfg.stream_so_cache)

    max_ep = _perguntar_epochs(cfg)
    if cfg.headless and max_ep <= 0:
        UI.warn("Headless com epochs=0 — usando 1 por seguranca."); max_ep = 1

    resultado = treinador.treinar(max_epochs=max_ep)
    if not resultado["historico"]:
        alertar("Sem historico de treino.", "Aviso")
        log_jsonl("erro", level="ERRO", motivo="sem_historico")
        if cfg.headless: sys.exit(1)
        return

    # F2.5 - reservatorio enriquecido
    reservoir = treinador.amostrador.get()
    if modo_completo and CAMPO_ID_DEFAULT in reservoir.columns and not reservoir.empty:
        UI.secao("F2.5", "Enriquecendo Reservoir com CNPJs")
        faltantes = reservoir[CAMPO_ID_DEFAULT].astype(str).nunique()
        UI.info(f"Reservoir: {len(reservoir)} linhas | ~{faltantes} CNPJs unicos")
        with UI.Spinner("Consultando APIs (cache + novos)..."):
            reservoir = enriquecer_cnpjs(reservoir, max_cnpjs=cfg.max_cnpj)
        reservoir = adicionar_features_avancadas(reservoir, col_txt, col_cat)
        reservoir = calcular_kpis_negocio(reservoir)
        reservoir["KPI_PadraoChar"] = aplicar_char_ngram(reservoir, col_txt, top_gramas)
        reservoir = calcular_risco(reservoir)
        treinador.amostrador.chunks = [reservoir]
        treinador.amostrador.n_rows = len(reservoir)
        n_ext = sum(1 for c in reservoir.columns if c.startswith("EXT_"))
        UI.ok(f"Reservoir enriquecido: {reservoir.shape[1]} colunas ({n_ext} EXT_*)")
        log_jsonl("reservoir_enriquecido",
                  n_linhas=len(reservoir), n_cnpjs=int(faltantes), n_ext=n_ext)
        UI.fim_secao()

    # F3
    UI.secao("F3", "Refinamento com SVC calibrado")
    metricas_final = resultado["ultima"]
    if cfg.sem_refinamento:
        UI.info("Refinamento desativado.")
        log_jsonl("refinamento_pulado")
    elif len(reservoir) > 200:
        try:
            if "KPI_PadraoChar" not in reservoir.columns:
                reservoir = adicionar_features_avancadas(reservoir, col_txt, col_cat)
                reservoir = calcular_kpis_negocio(reservoir)
                reservoir["KPI_PadraoChar"] = aplicar_char_ngram(
                    reservoir, col_txt, top_gramas)
                reservoir = calcular_risco(reservoir)
            X_res = preproc.transform(reservoir)
            ids_res = reservoir[col_id].astype(str).values
            mask_res = np.array([c in classes_map for c in ids_res])
            if mask_res.sum() > 200:
                y_res = np.array([classes_map[c] for c in ids_res[mask_res]])
                X_res_ok = X_res[mask_res]
                try:
                    strat = y_res if np.min(np.bincount(y_res)) >= 2 else None
                    X_tr, X_te, y_tr, y_te = train_test_split(
                        X_res_ok, y_res, test_size=0.2,
                        random_state=RANDOM_STATE, stratify=strat)
                    with UI.Spinner("Treinando modelo de refinamento..."):
                        modelo_ref = NT1Ensemble(alpha=1e-4, random_state=RANDOM_STATE)
                        modelo_ref.fit(X_tr, y_tr)
                        m_ref = avaliar_modelo(modelo_ref, X_te, y_te)
                        m_atual = avaliar_modelo(modelo, X_te, y_te)
                    chave = lambda m: m.get("top_5", m.get("accuracy", 0))
                    if chave(m_ref) > chave(m_atual):
                        UI.ok("Modelo refinado escolhido")
                        modelo = modelo_ref; metricas_final = m_ref
                    else:
                        UI.ok("Mantendo modelo streaming"); metricas_final = m_atual
                    UI.tabela(["Metrica", "Valor"],
                              [[k, v] for k, v in metricas_final.items()])
                    log_jsonl("refinamento_fim",
                              escolhido="refinado" if modelo is modelo_ref else "streaming",
                              metricas=metricas_final)
                except Exception as e:
                    UI.warn(f"refinamento: {e}")
                    log_jsonl("refinamento_erro", level="WARN", erro=str(e))
        except Exception as e:
            UI.warn(f"reservoir: {e}")
            log_jsonl("reservoir_erro", level="WARN", erro=str(e))
    UI.fim_secao()

    # F3.5
    if cfg.neural:
        try:
            if len(reservoir) > 200:
                if "KPI_PadraoChar" not in reservoir.columns:
                    UI.warn("Reservoir sem features — recomputando.")
                    reservoir = adicionar_features_avancadas(reservoir, col_txt, col_cat)
                    reservoir = calcular_kpis_negocio(reservoir)
                    reservoir["KPI_PadraoChar"] = aplicar_char_ngram(
                        reservoir, col_txt, top_gramas)
                    reservoir = calcular_risco(reservoir)
                n_ext = sum(1 for c in reservoir.columns if c.startswith("EXT_"))
                if modo_completo and n_ext == 0:
                    UI.warn("Rede neural treinando SEM features de CNPJ.")
                X_res = preproc.transform(reservoir)
                ids_res = reservoir[col_id].astype(str).values
                mask = np.array([c in classes_map for c in ids_res])
                if mask.sum() > 100:
                    y_nn = np.array([classes_map[c] for c in ids_res[mask]])
                    X_nn = X_res[mask]
                    trainer = NeuralTrainer(cfg, inv_classes)
                    met_nn = trainer.treinar(X_nn, y_nn)
                    if met_nn:
                        best_nn = met_nn.get("best_val_acc", 0.0)
                        best_atual = metricas_final.get(
                            "top_5", metricas_final.get("accuracy", 0))
                        UI.tabela(
                            ["Modelo", "Metrica Chave", "Valor"],
                            [["Ensemble", "top_5/acc", best_atual],
                             ["Rede Neural", "best_val_acc", best_nn]])
                        log_jsonl("neural_compare",
                                  ensemble=best_atual, neural=best_nn)
                        if best_nn > best_atual:
                            UI.ok("Rede Neural superou o ensemble — sera usada.")
                            modelo_nn = NeuralAdapter(trainer, np.arange(len(classes_map)))
                            modelo = modelo_nn
                            metricas_final = {**metricas_final,
                                              "neural_best_val_acc": best_nn}
                        else:
                            UI.info("Ensemble mantido (rede neural registrada no log).")
            else:
                UI.warn("Reservoir pequeno — rede neural pulada.")
        except Exception as e:
            UI.warn(f"rede neural: {e}")
            log_jsonl("neural_erro", level="WARN", erro=str(e))

    # F4
    UI.secao("F4", "Gerando resultados")
    partes = processar_e_escrever(
        cfg.pasta_dados, modelo, preproc, inv_classes, col_id,
        col_txt, col_cat, top_gramas, modo_completo)
    n_total = 0
    for p in partes:
        try: n_total += len(pd.read_parquet(p, engine="pyarrow"))
        except Exception: pass
    UI.ok(f"{len(partes)} partes | {n_total} linhas totais")
    log_jsonl("partes_ok", n_partes=len(partes), n_linhas=n_total)

    ts = time.strftime("%Y%m%d_%H%M%S")
    xlsx_path = os.path.join(cfg.pasta_saida, f"extract_v3_risco_{ts}.xlsx")
    if cfg.sem_xlsx:
        UI.info("XLSX desativado."); log_jsonl("xlsx_pulado")
    else:
        with UI.Spinner("Exportando XLSX top risco..."):
            exportar_top_risco(partes, xlsx_path, n=cfg.n_top_risco)
        log_jsonl("xlsx_ok", caminho=xlsx_path)

    top10_kpis = {}
    if partes:
        try:
            df_top = pd.concat([pd.read_parquet(p, engine="pyarrow")
                                for p in partes[:3]], ignore_index=True)
            df_top = df_top.sort_values("RiskScore", ascending=False).head(10)
            for c in COLUNAS_KPI:
                if c in df_top.columns:
                    top10_kpis[c] = round(float(pd.to_numeric(
                        df_top[c], errors="coerce").fillna(0).mean()), 4)
        except Exception: pass
    UI.fim_secao()
    log_jsonl("kpis_top10", **top10_kpis)

    # F5
    UI.secao("F5", "Persistencia e Historico")
    sessao = {
        "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
        "n_linhas": int(n_total), "n_classes": int(len(classes_map)),
        "modo": "completo" if modo_completo else "rapido",
        "headless": bool(cfg.headless),
        "neural": bool(cfg.neural),
        "stream_enriquecido": bool(cfg.stream_enriquecido),
        "colunas_selecionadas": {k: len(v) for k, v in selecionadas.items()},
        "metricas": metricas_final,
        "historico_epochs": resultado["historico"][-20:],
        "kpis_medio_top10": top10_kpis,
        "top_padroes_char": dict(list(top_gramas.items())[:20]),
    }
    hist = salvar_sessao(sessao)
    padrao = detectar_padrao(hist)
    salvar_padrao(padrao)
    UI.ok(f"Historico: {len(hist)} sessoes | "
          f"tendencia={padrao.get('tendencia','-')} | "
          f"pronto={padrao.get('pronto_para_palpite', False)}")
    log_jsonl("historico_ok", n_sessoes=len(hist),
              tendencia=padrao.get("tendencia"),
              pronto=padrao.get("pronto_para_palpite"))

    try:
        joblib.dump({
            "modelo": modelo, "preprocessador": preproc,
            "classes_map": classes_map, "inv_classes": inv_classes,
            "metricas": metricas_final, "colunas_selecionadas": selecionadas,
            "col_id": col_id, "top_gramas": top_gramas, "sessao": len(hist),
        }, cfg.modelo_path, compress=3)
        UI.ok(f"Modelo salvo em: {cfg.modelo_path}")
        log_jsonl("modelo_salvo", caminho=cfg.modelo_path)
    except Exception as e:
        UI.warn(f"joblib: {e}")
        log_jsonl("modelo_erro", level="WARN", erro=str(e))

    kpi_path = os.path.join(cfg.pasta_saida, f"extract_v3_kpis_{ts}.json")
    with open(kpi_path, "w", encoding="utf-8") as f:
        json.dump({
            "sessao": len(hist), "n_linhas": n_total,
            "n_classes": len(classes_map), "metricas": metricas_final,
            "colunas_selecionadas": selecionadas, "padrao": padrao,
            "top10_kpis": top10_kpis,
            "top_padroes_char": dict(list(top_gramas.items())[:30]),
        }, f, indent=2, ensure_ascii=False, default=str)
    UI.ok(f"KPIs salvos em: {kpi_path}")
    UI.fim_secao()
    log_jsonl("kpis_salvos", caminho=kpi_path)

    resumo = (f"Linhas: {n_total}\nClasses: {len(classes_map)}\n"
              f"Modo: {'Completo' if modo_completo else 'Rapido'}"
              f"{' (headless)' if cfg.headless else ''}\n"
              f"Rede Neural: {'on' if cfg.neural else 'off'}\n\n"
              f"Metricas:\n" +
              "\n".join(f"  {k}: {v}" for k, v in metricas_final.items()) +
              f"\n\nTendencia: {padrao.get('tendencia','-')}\n"
              f"Pronto: {padrao.get('pronto_para_palpite', False)}")

    if cfg.headless:
        UI.caixa(f"{APP_NOME} - sessao #{len(hist)}", resumo, cor=UI.GREEN)
        log_jsonl("sessao_fim", duracao_s=round(time.time() - t0, 2),
                  n_linhas=n_total, n_classes=len(classes_map),
                  metricas=metricas_final,
                  tendencia=padrao.get("tendencia"),
                  pronto=padrao.get("pronto_para_palpite"))
        UI.banner("FIM (headless)", f"Tempo total: {time.time() - t0:.2f}s")
        return

    alertar(f"{APP_NOME} - sessao #{len(hist)}\n\n{resumo}", "Extract V.3 Concluido")
    console_loop(modelo, preproc, inv_classes, col_txt, col_cat, top_gramas, col_id)
    log_jsonl("sessao_fim", duracao_s=round(time.time() - t0, 2))
    UI.banner("FIM", f"Tempo total: {time.time() - t0:.2f}s")


if __name__ == "__main__":
    main()
